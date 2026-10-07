import sqlite3
from pathlib import Path
from openpyxl import load_workbook


# ----------------------------------
# File locations
# ----------------------------------

input_file = "output/session_summary.xlsx"
database_file = "output/training.db"


# ----------------------------------
# Connect to database
# ----------------------------------

conn = sqlite3.connect(database_file)
cursor = conn.cursor()


# ----------------------------------
# Create Activities table
# ----------------------------------

cursor.execute("""
CREATE TABLE IF NOT EXISTS activities (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    activity_date TEXT,
    session_name TEXT,
    session_type TEXT,
    sport TEXT,
    sub_sport TEXT,
    environment TEXT,
    distance_km REAL,
    time_hms TEXT,
    average_hr REAL,
    average_speed REAL,
    average_pace TEXT,
    ascent_m REAL,
    calories REAL,
    notes TEXT,
    rpe REAL
)
""")


# ----------------------------------
# Open Session Summary workbook
# ----------------------------------

workbook = load_workbook(input_file)
sheet = workbook["Session Summary"]

headers = [cell.value for cell in sheet[1]]

date_col = headers.index("Date") + 1
session_name_col = headers.index("Session Name") + 1
session_type_col = headers.index("Session Type") + 1
sport_col = headers.index("Sport") + 1
sub_sport_col = headers.index("Sub Sport") + 1
environment_col = headers.index("Environment") + 1
distance_col = headers.index("Distance (km)") + 1
time_col = headers.index("Time (hh:mm:ss)") + 1
hr_col = headers.index("Average Heart Rate (bpm)") + 1
speed_col = headers.index("Average Speed (m/s)") + 1
pace_col = headers.index("Average Pace (min/km)") + 1
ascent_col = headers.index("Total Ascent (m)") + 1
calories_col = headers.index("Calories") + 1
notes_col = headers.index("Notes") + 1
rpe_col = headers.index("RPE") + 1


# ----------------------------------
# Clear existing records
# ----------------------------------

cursor.execute("DELETE FROM activities")


# ----------------------------------
# Insert rows
# ----------------------------------

rows_imported = 0

for row in range(2, sheet.max_row + 1):

    cursor.execute("""
    INSERT INTO activities (
        activity_date,
        session_name,
        session_type,
        sport,
        sub_sport,
        environment,
        distance_km,
        time_hms,
        average_hr,
        average_speed,
        average_pace,
        ascent_m,
        calories,
        notes,
        rpe
    )
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (

        sheet.cell(row, date_col).value,
        sheet.cell(row, session_name_col).value,
        sheet.cell(row, session_type_col).value,
        sheet.cell(row, sport_col).value,
        sheet.cell(row, sub_sport_col).value,
        sheet.cell(row, environment_col).value,
        sheet.cell(row, distance_col).value,
        sheet.cell(row, time_col).value,
        sheet.cell(row, hr_col).value,
        sheet.cell(row, speed_col).value,
        sheet.cell(row, pace_col).value,
        sheet.cell(row, ascent_col).value,
        sheet.cell(row, calories_col).value,
        sheet.cell(row, notes_col).value,
        sheet.cell(row, rpe_col).value

    ))

    rows_imported += 1


conn.commit()

print()
print("=" * 40)
print("DATABASE IMPORT COMPLETE")
print("=" * 40)
print(f"Rows imported: {rows_imported}")
print(f"Database: {database_file}")

conn.close()