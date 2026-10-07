from openpyxl import load_workbook, Workbook

input_file = "output/tempo_work_block_summary.xlsx"
output_file = "output/tempo_work_block_trend.xlsx"

# Open tempo work block summary file
workbook = load_workbook(input_file)
sheet = workbook["Tempo Work Summary"]

# Read headers
headers = [cell.value for cell in sheet[1]]

date_col = headers.index("Date") + 1
session_name_col = headers.index("Session Name") + 1
work_blocks_col = headers.index("Work Blocks") + 1
work_laps_col = headers.index("Work Laps") + 1
work_distance_col = headers.index("Work Distance (km)") + 1
work_time_col = headers.index("Work Time") + 1
work_pace_col = headers.index("Work Pace") + 1
work_avg_hr_col = headers.index("Work Avg HR") + 1
work_max_hr_col = headers.index("Work Max HR") + 1
work_ascent_col = headers.index("Work Ascent (m)") + 1
work_efficiency_col = headers.index("Work Efficiency Score") + 1


# Store rows
tempo_rows = []

for row in range(2, sheet.max_row + 1):
    date = sheet.cell(row=row, column=date_col).value
    session_name = sheet.cell(row=row, column=session_name_col).value
    work_blocks = sheet.cell(row=row, column=work_blocks_col).value
    work_laps = sheet.cell(row=row, column=work_laps_col).value
    work_distance = sheet.cell(row=row, column=work_distance_col).value
    work_time = sheet.cell(row=row, column=work_time_col).value
    work_pace = sheet.cell(row=row, column=work_pace_col).value
    work_avg_hr = sheet.cell(row=row, column=work_avg_hr_col).value
    work_max_hr = sheet.cell(row=row, column=work_max_hr_col).value
    work_ascent = sheet.cell(row=row, column=work_ascent_col).value
    work_efficiency = sheet.cell(row=row, column=work_efficiency_col).value

    if date is None:
        continue

    tempo_rows.append({
        "date": date,
        "session_name": session_name,
        "work_blocks": work_blocks,
        "work_laps": work_laps,
        "work_distance": work_distance,
        "work_time": work_time,
        "work_pace": work_pace,
        "work_avg_hr": work_avg_hr,
        "work_max_hr": work_max_hr,
        "work_ascent": work_ascent,
        "work_efficiency": work_efficiency
    })


# Sort by date
tempo_rows = sorted(tempo_rows, key=lambda x: x["date"])

# Calculate change vs previous and 4-session rolling average
previous_efficiency = None
efficiency_values = []

for row in tempo_rows:
    efficiency = row["work_efficiency"]

    change_vs_previous = None
    if previous_efficiency is not None and efficiency is not None:
        change_vs_previous = efficiency - previous_efficiency

    if efficiency is not None:
        efficiency_values.append(efficiency)
        previous_efficiency = efficiency

    rolling_average = None
    if efficiency_values:
        last_four_values = efficiency_values[-4:]
        rolling_average = sum(last_four_values) / len(last_four_values)

    row["change_vs_previous"] = change_vs_previous
    row["rolling_average"] = rolling_average


# Create output workbook
output_workbook = Workbook()
trend_sheet = output_workbook.active
trend_sheet.title = "Tempo Work Trend"

trend_sheet.append([
    "Date",
    "Session Name",
    "Work Blocks",
    "Work Laps",
    "Work Distance (km)",
    "Work Time",
    "Work Pace",
    "Work Avg HR",
    "Work Max HR",
    "Work Ascent (m)",
    "Work Efficiency Score",
    "Change vs Previous",
    "4-Session Rolling Avg Efficiency"
])

for row in tempo_rows:
    trend_sheet.append([
        row["date"],
        row["session_name"],
        row["work_blocks"],
        row["work_laps"],
        row["work_distance"],
        row["work_time"],
        row["work_pace"],
        row["work_avg_hr"],
        row["work_max_hr"],
        row["work_ascent"],
        row["work_efficiency"],
        round(row["change_vs_previous"], 2) if row["change_vs_previous"] is not None else None,
        round(row["rolling_average"], 2) if row["rolling_average"] is not None else None
    ])


# Set column widths
column_widths = {
    "A": 14,
    "B": 18,
    "C": 14,
    "D": 12,
    "E": 18,
    "F": 14,
    "G": 14,
    "H": 14,
    "I": 14,
    "J": 16,
    "K": 20,
    "L": 20,
    "M": 30
}

for col, width in column_widths.items():
    trend_sheet.column_dimensions[col].width = width


output_workbook.save(output_file)

print(f"Tempo work block trend saved to: {output_file}")