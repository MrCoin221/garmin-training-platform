from openpyxl import load_workbook, Workbook

input_file = "output/lap_summary.xlsx"
output_file = "output/tempo_work_block_summary.xlsx"

TEMPO_WORK_PACE_CUTOFF = "04:35"
MIN_WORK_LAP_DISTANCE_KM = 0.5


def hms_to_seconds(time_value):
    if time_value is None:
        return 0

    parts = str(time_value).split(":")
    if len(parts) != 3:
        return 0

    return int(parts[0]) * 3600 + int(parts[1]) * 60 + int(parts[2])


def pace_to_seconds(pace_value):
    if pace_value is None:
        return None

    parts = str(pace_value).split(":")
    if len(parts) != 2:
        return None

    return int(parts[0]) * 60 + int(parts[1])


def seconds_to_hms(total_seconds):
    total_seconds = int(total_seconds)
    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    seconds = total_seconds % 60
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def seconds_to_pace(seconds_per_km):
    if seconds_per_km is None:
        return None

    minutes = int(seconds_per_km // 60)
    seconds = int(seconds_per_km % 60)
    return f"{minutes:02d}:{seconds:02d}"


def calculate_efficiency(distance_km, time_seconds, avg_hr):
    if distance_km <= 0 or time_seconds <= 0 or avg_hr is None or avg_hr <= 0:
        return None

    distance_m = distance_km * 1000
    avg_speed_mps = distance_m / time_seconds
    return (avg_speed_mps / avg_hr) * 1000


# Open lap summary workbook
workbook = load_workbook(input_file)
sheet = workbook["Lap Summary"]

headers = [cell.value for cell in sheet[1]]

date_col = headers.index("Date") + 1
session_name_col = headers.index("Session Name") + 1
session_type_col = headers.index("Session Type") + 1
environment_col = headers.index("Environment") + 1
lap_number_col = headers.index("Lap Number") + 1
lap_distance_col = headers.index("Lap Distance (km)") + 1
lap_time_col = headers.index("Lap Time (hh:mm:ss)") + 1
lap_pace_col = headers.index("Lap Pace (min/km)") + 1
lap_avg_hr_col = headers.index("Lap Avg Heart Rate (bpm)") + 1
lap_max_hr_col = headers.index("Lap Max Heart Rate (bpm)") + 1
lap_ascent_col = headers.index("Lap Ascent (m)") + 1

cutoff_seconds = pace_to_seconds(TEMPO_WORK_PACE_CUTOFF)

work_laps = []

# Identify tempo work laps
for row in range(2, sheet.max_row + 1):
    date = sheet.cell(row=row, column=date_col).value
    session_name = sheet.cell(row=row, column=session_name_col).value
    session_type = sheet.cell(row=row, column=session_type_col).value
    environment = sheet.cell(row=row, column=environment_col).value
    lap_number = sheet.cell(row=row, column=lap_number_col).value
    lap_distance = sheet.cell(row=row, column=lap_distance_col).value
    lap_time = sheet.cell(row=row, column=lap_time_col).value
    lap_pace = sheet.cell(row=row, column=lap_pace_col).value
    lap_avg_hr = sheet.cell(row=row, column=lap_avg_hr_col).value
    lap_max_hr = sheet.cell(row=row, column=lap_max_hr_col).value
    lap_ascent = sheet.cell(row=row, column=lap_ascent_col).value

    if session_type != "Tempo Run":
        continue

    if environment != "Outdoor":
        continue

    if lap_distance is None or lap_distance < MIN_WORK_LAP_DISTANCE_KM:
        continue

    lap_pace_seconds = pace_to_seconds(lap_pace)

    if lap_pace_seconds is None:
        continue

    if lap_pace_seconds > cutoff_seconds:
        continue

    lap_time_seconds = hms_to_seconds(lap_time)

    work_laps.append({
        "date": date,
        "session_name": session_name,
        "lap_number": lap_number,
        "distance": lap_distance,
        "time_seconds": lap_time_seconds,
        "pace": lap_pace,
        "avg_hr": lap_avg_hr,
        "max_hr": lap_max_hr,
        "ascent": lap_ascent if lap_ascent is not None else 0
    })


# Add work block numbers based on consecutive lap numbers within same session
previous_key = None
previous_lap_number = None
current_block_number = 0

for lap in work_laps:
    current_key = (lap["date"], lap["session_name"])

    if current_key != previous_key:
        current_block_number = 1
    elif previous_lap_number is not None and lap["lap_number"] == previous_lap_number + 1:
        pass
    else:
        current_block_number += 1

    lap["work_block"] = current_block_number

    previous_key = current_key
    previous_lap_number = lap["lap_number"]


# Summarise by tempo session
session_summary = {}

for lap in work_laps:
    key = (lap["date"], lap["session_name"])

    if key not in session_summary:
        session_summary[key] = {
            "date": lap["date"],
            "session_name": lap["session_name"],
            "work_blocks": set(),
            "work_laps": 0,
            "distance": 0,
            "time_seconds": 0,
            "weighted_hr_total": 0,
            "hr_time_weight": 0,
            "max_hr": None,
            "ascent": 0
        }

    session_summary[key]["work_blocks"].add(lap["work_block"])
    session_summary[key]["work_laps"] += 1
    session_summary[key]["distance"] += lap["distance"]
    session_summary[key]["time_seconds"] += lap["time_seconds"]
    session_summary[key]["ascent"] += lap["ascent"]

    if lap["avg_hr"] is not None:
        session_summary[key]["weighted_hr_total"] += lap["avg_hr"] * lap["time_seconds"]
        session_summary[key]["hr_time_weight"] += lap["time_seconds"]

    if lap["max_hr"] is not None:
        if session_summary[key]["max_hr"] is None or lap["max_hr"] > session_summary[key]["max_hr"]:
            session_summary[key]["max_hr"] = lap["max_hr"]


# Create output workbook
output_workbook = Workbook()

# Sheet 1: Work Laps
work_laps_sheet = output_workbook.active
work_laps_sheet.title = "Tempo Work Laps"

work_laps_sheet.append([
    "Date",
    "Session Name",
    "Work Block",
    "Lap Number",
    "Lap Distance (km)",
    "Lap Time",
    "Lap Pace",
    "Lap Avg HR",
    "Lap Max HR",
    "Lap Ascent (m)"
])

for lap in work_laps:
    work_laps_sheet.append([
        lap["date"],
        lap["session_name"],
        lap["work_block"],
        lap["lap_number"],
        lap["distance"],
        seconds_to_hms(lap["time_seconds"]),
        lap["pace"],
        lap["avg_hr"],
        lap["max_hr"],
        lap["ascent"]
    ])


# Sheet 2: Session Summary
summary_sheet = output_workbook.create_sheet("Tempo Work Summary")

summary_sheet.append([
    "Date",
    "Session Name",
    "Work Blocks",
    "Work Laps",
    "Work Distance (km)",
    "Work Time",
    "Work Pace",
    "Work Avg HR",
    "Work Max HR",
    "Work Ascent (m)",
    "Work Efficiency Score"
])

for key in sorted(session_summary.keys()):
    data = session_summary[key]

    avg_hr = None
    if data["hr_time_weight"] > 0:
        avg_hr = data["weighted_hr_total"] / data["hr_time_weight"]

    work_pace = None
    if data["distance"] > 0:
        work_pace_seconds = data["time_seconds"] / data["distance"]
        work_pace = seconds_to_pace(work_pace_seconds)

    efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        avg_hr
    )

    summary_sheet.append([
        data["date"],
        data["session_name"],
        len(data["work_blocks"]),
        data["work_laps"],
        round(data["distance"], 2),
        seconds_to_hms(data["time_seconds"]),
        work_pace,
        round(avg_hr, 1) if avg_hr else None,
        data["max_hr"],
        round(data["ascent"], 0),
        round(efficiency, 2) if efficiency else None
    ])


# Set column widths
for worksheet in [work_laps_sheet, summary_sheet]:
    for col in range(1, worksheet.max_column + 1):
        worksheet.column_dimensions[worksheet.cell(row=1, column=col).column_letter].width = 18

output_workbook.save(output_file)

print(f"Tempo work block summary saved to: {output_file}")