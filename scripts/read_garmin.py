from fitparse import FitFile
import os
from openpyxl import Workbook, load_workbook

data_folder = "data"
output_folder = "output"
output_file = os.path.join(output_folder, "session_summary.xlsx")

files = os.listdir(data_folder)
fit_files = sorted([file for file in files if file.endswith(".fit")])


def classify_session_type(session_name):
    """Determine session type from file name."""
    name_lower = session_name.lower()

    if "brick_run" in name_lower:
        return "Brick Run"
    elif "easy_run" in name_lower:
        return "Easy Run"
    elif "interval_run" in name_lower:
        return "Interval Run"
    elif "long_run" in name_lower:
        return "Long Run"
    elif "park_run" in name_lower:
        return "Park Run"
    elif "recovery_run" in name_lower:
        return "Recovery Run"
    elif "race_day" in name_lower:
        return "Race Day"
    elif "tempo_run" in name_lower:
        return "Tempo Run"
    elif "threshold_1" in name_lower:
        return "Threshold 1"
    elif "threshold_2" in name_lower:
        return "Threshold 2"
    else:
        return "Other"


def load_existing_manual_fields():
    """
    Read existing Notes and RPE from session_summary.xlsx.

    Returns a dictionary:
    {
        source_file: {
            "notes": "...",
            "rpe": "..."
        }
    }
    """
    manual_fields = {}

    if not os.path.exists(output_file):
        return manual_fields

    workbook = load_workbook(output_file)
    sheet = workbook["Session Summary"]

    headers = [cell.value for cell in sheet[1]]

    notes_col = headers.index("Notes") + 1 if "Notes" in headers else None
    rpe_col = headers.index("RPE") + 1 if "RPE" in headers else None

    source_file_col = headers.index("Source File") + 1 if "Source File" in headers else None
    date_col = headers.index("Date") + 1 if "Date" in headers else None
    session_name_col = headers.index("Session Name") + 1 if "Session Name" in headers else None

    for row in range(2, sheet.max_row + 1):
        source_file = None

        # Preferred method: use Source File if it exists
        if source_file_col is not None:
            source_file = sheet.cell(row=row, column=source_file_col).value

        # Fallback for old versions of the workbook that did not yet have Source File
        if source_file is None and date_col is not None and session_name_col is not None:
            date = sheet.cell(row=row, column=date_col).value
            session_name = sheet.cell(row=row, column=session_name_col).value

            if date is not None and session_name is not None:
                source_file = f"{date}_{session_name}.fit"

        if source_file is None:
            continue

        notes = sheet.cell(row=row, column=notes_col).value if notes_col is not None else ""
        rpe = sheet.cell(row=row, column=rpe_col).value if rpe_col is not None else ""

        manual_fields[source_file] = {
            "notes": notes if notes is not None else "",
            "rpe": rpe if rpe is not None else ""
        }

    return manual_fields


if not fit_files:
    print("No .fit files found in the data folder.")

else:
    existing_manual_fields = load_existing_manual_fields()
    rows = []

    for file_name in fit_files:
        file_path = os.path.join(data_folder, file_name)
        print(f"\nReading file: {file_name}")

        # Split file name into date and session name
        base_name = file_name.replace(".fit", "")
        parts = base_name.split("_", 1)

        if len(parts) == 2:
            date = parts[0]
            session_name = parts[1]
        else:
            date = ""
            session_name = base_name

        session_type = classify_session_type(session_name)

        # Open FIT file
        fitfile = FitFile(file_path)

        # Session-level summary fields
        total_distance_m = None
        total_timer_time_s = None
        avg_heart_rate_bpm = None
        avg_speed_mps = None
        total_ascent_m = None
        total_calories = None
        sport = None
        sub_sport = None

        # Read session data from FIT file
        for session in fitfile.get_messages("session"):
            for data in session:
                if data.name == "total_distance":
                    total_distance_m = data.value
                elif data.name == "total_timer_time":
                    total_timer_time_s = data.value
                elif data.name == "avg_heart_rate":
                    avg_heart_rate_bpm = data.value
                elif data.name == "enhanced_avg_speed":
                    avg_speed_mps = data.value
                elif data.name == "total_ascent":
                    total_ascent_m = data.value
                elif data.name == "total_calories":
                    total_calories = data.value
                elif data.name == "sport":
                    sport = data.value
                elif data.name == "sub_sport":
                    sub_sport = data.value

        # Determine environment
        sub_sport_text = str(sub_sport).lower() if sub_sport is not None else ""
        name_lower = session_name.lower()

        if "treadmill" in sub_sport_text:
            environment = "Treadmill"
        elif "indoor" in sub_sport_text:
            environment = "Indoor"
        elif "treadmill" in name_lower:
            environment = "Treadmill"
        elif total_ascent_m is None:
            environment = "Unknown"
        else:
            environment = "Outdoor"

        # Clean up sub sport for Excel output
        if sub_sport == "generic":
            sub_sport_for_excel = ""
        else:
            sub_sport_for_excel = sub_sport

        # Distance in km
        distance_km = None
        if total_distance_m is not None:
            distance_km = round(total_distance_m / 1000, 2)

        # Time in hh:mm:ss
        time_hms = None
        if total_timer_time_s is not None:
            total_seconds = int(total_timer_time_s)
            hours = total_seconds // 3600
            minutes = (total_seconds % 3600) // 60
            seconds = total_seconds % 60
            time_hms = f"{hours:02d}:{minutes:02d}:{seconds:02d}"

        # Pace from speed
        avg_pace = None
        if avg_speed_mps is not None and avg_speed_mps > 0:
            avg_pace_sec_per_km = 1000 / avg_speed_mps
            pace_min = int(avg_pace_sec_per_km // 60)
            pace_sec = int(avg_pace_sec_per_km % 60)
            avg_pace = f"{pace_min:02d}:{pace_sec:02d}"

        # Preserve existing manual Notes and RPE
        existing_manual = existing_manual_fields.get(file_name, {})
        notes = existing_manual.get("notes", "")
        rpe = existing_manual.get("rpe", "")

        # Print summary
        print(f"Date: {date}")
        print(f"Session Name: {session_name}")
        print(f"Session Type: {session_type}")
        print(f"Sport: {sport}")
        print(f"Sub Sport: {sub_sport_for_excel}")
        print(f"Environment: {environment}")
        print(f"Distance: {distance_km} km")
        print(f"Time: {time_hms}")
        print(f"Average Heart Rate: {avg_heart_rate_bpm} bpm")
        print(
            f"Average Speed: {avg_speed_mps:.3f} m/s"
            if avg_speed_mps is not None
            else "Average Speed: None"
        )
        print(f"Average Pace: {avg_pace} min/km")
        print(f"Total Ascent: {total_ascent_m} m")
        print(f"Calories: {total_calories}")
        print(f"Notes preserved: {'Yes' if notes else 'No'}")
        print(f"RPE preserved: {'Yes' if rpe else 'No'}")

        rows.append([
            file_name,
            date,
            session_name,
            session_type,
            sport,
            sub_sport_for_excel,
            environment,
            distance_km,
            time_hms,
            avg_heart_rate_bpm,
            round(avg_speed_mps, 3) if avg_speed_mps is not None else None,
            avg_pace,
            total_ascent_m,
            total_calories,
            notes,
            rpe
        ])

    # Create Excel workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Session Summary"

    # Header row
    ws.append([
        "Source File",
        "Date",
        "Session Name",
        "Session Type",
        "Sport",
        "Sub Sport",
        "Environment",
        "Distance (km)",
        "Time (hh:mm:ss)",
        "Average Heart Rate (bpm)",
        "Average Speed (m/s)",
        "Average Pace (min/km)",
        "Total Ascent (m)",
        "Calories",
        "Notes",
        "RPE"
    ])

    # Data rows
    for row in rows:
        ws.append(row)

    # Set column widths
    column_widths = {
        "A": 30,
        "B": 14,
        "C": 24,
        "D": 18,
        "E": 12,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 16,
        "J": 24,
        "K": 20,
        "L": 22,
        "M": 18,
        "N": 12,
        "O": 40,
        "P": 8
    }

    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    wb.save(output_file)

    print(f"\nSession summaries saved to: {output_file}")