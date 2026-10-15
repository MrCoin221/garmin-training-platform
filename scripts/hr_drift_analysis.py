import os
import re
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill


LAP_FILE = "output/lap_summary.xlsx"
PLAN_FILE = "data/training_plan.xlsx"
TREND_FILE = "output/trend_summary.xlsx"

LAP_SHEET = "Lap Summary"
PLAN_SHEET = "Planned Sessions"

DRIFT_SHEETS = {
    "Easy Run": "Easy HR Drift",
    "Long Run": "Long HR Drift"
}

WORK_SHEET = "Running Work Segments"


def safe_float(value, default=0.0):
    try:
        if value is None:
            return default

        return float(value)

    except (TypeError, ValueError):
        return default


def time_to_seconds(value):
    if value is None:
        return 0.0

    parts = str(value).split(":")

    if len(parts) != 3:
        return 0.0

    try:
        return (
            int(parts[0]) * 3600
            + int(parts[1]) * 60
            + float(parts[2])
        )

    except (TypeError, ValueError):
        return 0.0


def seconds_to_excel_time(seconds):
    return seconds / 86400


def parse_distance_segment(segment_text):
    """
    Examples:
        blank -> None
        Full -> None
        0-12 -> (0.0, 12.0)
        12-16 -> (12.0, 16.0)
    """

    if segment_text is None:
        return None

    cleaned = str(segment_text).strip()

    if not cleaned:
        return None

    if cleaned.lower() in {
        "full",
        "full run",
        "full session"
    }:
        return None

    match = re.fullmatch(
        r"\s*(\d+(?:\.\d+)?)\s*-\s*"
        r"(\d+(?:\.\d+)?)\s*",
        cleaned
    )

    if not match:
        raise ValueError(
            f"Invalid segment: {segment_text}. "
            "Use a format such as 0-12."
        )

    start_km = float(match.group(1))
    end_km = float(match.group(2))

    if end_km <= start_km:
        raise ValueError(
            f"Segment end must exceed start: "
            f"{segment_text}"
        )

    return start_km, end_km


def parse_work_segments(segment_text):
    """
    Example:
        12-16:Controlled;16-17:Recovery
    """

    if segment_text is None:
        return []

    cleaned = str(segment_text).strip()

    if not cleaned:
        return []

    parsed_segments = []

    for segment_item in cleaned.split(";"):
        segment_item = segment_item.strip()

        if not segment_item:
            continue

        if ":" in segment_item:
            distance_part, name = segment_item.split(
                ":",
                1
            )
            name = name.strip()
        else:
            distance_part = segment_item
            name = "Work Segment"

        boundaries = parse_distance_segment(
            distance_part
        )

        if boundaries is None:
            continue

        start_km, end_km = boundaries

        parsed_segments.append({
            "name": name,
            "start_km": start_km,
            "end_km": end_km
        })

    return parsed_segments


def normalise_date(value):
    if value is None:
        return ""

    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")

    return str(value).split(" ")[0]


def get_plan_value(
    plan_record,
    *possible_headers
):
    """
    Retrieve a training-plan value while ignoring
    differences in spaces and capitalisation.
    """

    if not plan_record:
        return None

    normalised_record = {
        str(key).strip().lower(): value
        for key, value in plan_record.items()
        if key is not None
    }

    for header in possible_headers:
        normalised_header = (
            str(header)
            .strip()
            .lower()
        )

        if normalised_header in normalised_record:
            return normalised_record[
                normalised_header
            ]

    return None


def read_worksheet_records(
    worksheet,
    required_columns=None
):
    headers = [
        (
            str(cell.value).strip()
            if cell.value is not None
            else None
        )
        for cell in worksheet[1]
    ]

    column_positions = {
        header: index
        for index, header in enumerate(headers)
        if header is not None
    }

    if required_columns:
        missing = [
            column
            for column in required_columns
            if column not in column_positions
        ]

        if missing:
            raise ValueError(
                f"Missing columns in "
                f"{worksheet.title}: {missing}"
            )

    records = []

    for values in worksheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        record = {}

        for header, index in column_positions.items():
            record[header] = values[index]

        records.append(record)

    return records


def get_segment_from_laps(
    laps,
    start_km,
    end_km
):
    """
    Extract a distance segment while proportionally
    splitting laps crossed by a boundary.
    """

    total_distance = 0.0
    total_time = 0.0
    weighted_hr = 0.0
    weighted_hr_time = 0.0
    max_hr = 0.0
    ascent = 0.0

    current_distance = 0.0

    for lap in laps:
        lap_distance = safe_float(
            lap["Lap Distance (km)"]
        )

        lap_time = time_to_seconds(
            lap["Lap Time (hh:mm:ss)"]
        )

        lap_avg_hr = safe_float(
            lap["Lap Avg Heart Rate (bpm)"]
        )

        lap_max_hr = safe_float(
            lap.get("Lap Max Heart Rate (bpm)")
        )

        lap_ascent = safe_float(
            lap["Lap Ascent (m)"]
        )

        lap_start = current_distance
        lap_end = current_distance + lap_distance

        overlap_start = max(
            lap_start,
            start_km
        )

        overlap_end = min(
            lap_end,
            end_km
        )

        overlap_distance = max(
            0.0,
            overlap_end - overlap_start
        )

        if (
            overlap_distance > 0
            and lap_distance > 0
        ):
            fraction = (
                overlap_distance
                / lap_distance
            )

            allocated_time = (
                lap_time * fraction
            )

            total_distance += overlap_distance
            total_time += allocated_time
            ascent += lap_ascent * fraction

            if lap_avg_hr > 0:
                weighted_hr += (
                    lap_avg_hr
                    * allocated_time
                )

                weighted_hr_time += (
                    allocated_time
                )

            max_hr = max(
                max_hr,
                lap_max_hr
            )

        current_distance = lap_end

        if current_distance >= end_km:
            break

    average_hr = (
        weighted_hr / weighted_hr_time
        if weighted_hr_time > 0
        else None
    )

    pace_seconds_per_km = (
        total_time / total_distance
        if total_distance > 0
        else None
    )

    speed_kmh = (
        total_distance
        / (total_time / 3600)
        if total_time > 0
        else None
    )

    efficiency = (
        speed_kmh / average_hr
        if speed_kmh
        and average_hr
        and average_hr > 0
        else None
    )

    return {
        "distance_km": total_distance,
        "time_seconds": total_time,
        "pace_seconds_per_km":
            pace_seconds_per_km,
        "average_hr": average_hr,
        "max_hr": max_hr,
        "ascent_m": ascent,
        "speed_kmh": speed_kmh,
        "efficiency": efficiency
    }


def calculate_drift(
    laps,
    segment_start,
    segment_end
):
    segment_distance = (
        segment_end - segment_start
    )

    midpoint = (
        segment_start
        + segment_distance / 2
    )

    first_half = get_segment_from_laps(
        laps,
        segment_start,
        midpoint
    )

    second_half = get_segment_from_laps(
        laps,
        midpoint,
        segment_end
    )

    first_efficiency = first_half[
        "efficiency"
    ]

    second_efficiency = second_half[
        "efficiency"
    ]

    if (
        first_efficiency is None
        or second_efficiency is None
        or first_efficiency == 0
    ):
        decoupling = None
    else:
        decoupling = (
            1
            - (
                second_efficiency
                / first_efficiency
            )
        )

    first_hr = first_half["average_hr"]
    second_hr = second_half["average_hr"]

    hr_drift = (
        second_hr / first_hr - 1
        if first_hr
        and second_hr
        else None
    )

    first_pace = first_half[
        "pace_seconds_per_km"
    ]

    second_pace = second_half[
        "pace_seconds_per_km"
    ]

    pace_change = (
        second_pace / first_pace - 1
        if first_pace
        and second_pace
        else None
    )

    return {
        "segment_start": segment_start,
        "segment_end": segment_end,
        "midpoint": midpoint,
        "first_half": first_half,
        "second_half": second_half,
        "hr_drift": hr_drift,
        "pace_change": pace_change,
        "decoupling": decoupling
    }


def format_sheet_header(worksheet):
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )


def set_column_widths(worksheet):
    widths = {
        "A": 13,
        "B": 19,
        "C": 18,
        "D": 14,
        "E": 14,
        "F": 14,
        "G": 14,
        "H": 14,
        "I": 14,
        "J": 14,
        "K": 14,
        "L": 14,
        "M": 14,
        "N": 14,
        "O": 14,
        "P": 14,
        "Q": 14,
        "R": 16,
        "S": 20
    }

    for column, width in widths.items():
        worksheet.column_dimensions[
            column
        ].width = width


def write_drift_sheet(
    workbook,
    sheet_name,
    rows
):
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]

    worksheet = workbook.create_sheet(
        sheet_name
    )

    headers = [
        "Date",
        "Session Name",
        "Analysis Segment",
        "Segment Distance (km)",
        "First Half Range",
        "First Half Pace",
        "First Half Avg HR",
        "First Half Ascent",
        "Second Half Range",
        "Second Half Pace",
        "Second Half Avg HR",
        "Second Half Ascent",
        "HR Drift %",
        "Pace Change %",
        "Aerobic Decoupling %",
        "Ascent Difference",
        "Analysis Source",
        "Context Flag",
        "Notes"
    ]

    worksheet.append(headers)

    for drift_row in rows:
        worksheet.append(drift_row)

    format_sheet_header(worksheet)
    set_column_widths(worksheet)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for row_number in range(
        2,
        worksheet.max_row + 1
    ):
        worksheet.cell(
            row_number,
            1
        ).number_format = "yyyy-mm-dd"

        worksheet.cell(
            row_number,
            6
        ).number_format = "mm:ss"

        worksheet.cell(
            row_number,
            10
        ).number_format = "mm:ss"

        for column in [13, 14, 15]:
            worksheet.cell(
                row_number,
                column
            ).number_format = "0.0%"

        for column in [
            4, 7, 8, 11, 12, 16
        ]:
            worksheet.cell(
                row_number,
                column
            ).number_format = "0.0"

    if worksheet.max_row >= 2:
        worksheet.conditional_formatting.add(
            f"O2:O{worksheet.max_row}",
            ColorScaleRule(
                start_type="num",
                start_value=0,
                start_color="C6EFCE",
                mid_type="num",
                mid_value=0.05,
                mid_color="FFEB9C",
                end_type="num",
                end_value=0.10,
                end_color="FFC7CE"
            )
        )


def write_work_segments_sheet(
    workbook,
    rows
):
    if WORK_SHEET in workbook.sheetnames:
        del workbook[WORK_SHEET]

    worksheet = workbook.create_sheet(
        WORK_SHEET
    )

    headers = [
        "Date",
        "Session Name",
        "Block Name",
        "Start Distance",
        "End Distance",
        "Block Distance",
        "Block Time",
        "Average Pace",
        "Average HR",
        "Maximum HR",
        "Ascent",
        "Efficiency",
        "Notes"
    ]

    worksheet.append(headers)

    for work_row in rows:
        worksheet.append(work_row)

    format_sheet_header(worksheet)
    set_column_widths(worksheet)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for row_number in range(
        2,
        worksheet.max_row + 1
    ):
        worksheet.cell(
            row_number,
            1
        ).number_format = "yyyy-mm-dd"

        worksheet.cell(
            row_number,
            7
        ).number_format = "mm:ss"

        worksheet.cell(
            row_number,
            8
        ).number_format = "mm:ss"

        for column in [
            4, 5, 6, 9, 10, 11, 12
        ]:
            worksheet.cell(
                row_number,
                column
            ).number_format = "0.00"


def main():
    if not os.path.exists(LAP_FILE):
        raise FileNotFoundError(
            f"Missing: {LAP_FILE}"
        )

    if not os.path.exists(TREND_FILE):
        raise FileNotFoundError(
            f"Missing: {TREND_FILE}"
        )

    lap_workbook = load_workbook(
        LAP_FILE,
        data_only=True
    )

    lap_worksheet = lap_workbook[
        LAP_SHEET
    ]

    lap_records = read_worksheet_records(
        lap_worksheet,
        required_columns=[
            "Date",
            "Session Name",
            "Session Type",
            "Environment",
            "Lap Number",
            "Lap Distance (km)",
            "Lap Time (hh:mm:ss)",
            "Lap Avg Heart Rate (bpm)",
            "Lap Ascent (m)"
        ]
    )

    plan_by_date = defaultdict(list)

    if os.path.exists(PLAN_FILE):
        plan_workbook = load_workbook(
            PLAN_FILE,
            data_only=True
        )

        plan_worksheet = plan_workbook[
            PLAN_SHEET
        ]

        plan_records = read_worksheet_records(
            plan_worksheet
        )

        for plan_record in plan_records:
            plan_date = normalise_date(
                plan_record.get("Date")
            )

            if plan_date:
                plan_by_date[plan_date].append(
                    plan_record
                )

        sessions = defaultdict(list)

    for lap_record in lap_records:

        # ------------------------------------------
        # Standardise the session type
        # ------------------------------------------

        session_type_raw = str(
            lap_record.get(
                "Session Type"
            )
            or ""
        ).strip()

        session_type_key = (
            session_type_raw
            .replace("_", " ")
            .lower()
        )

        if session_type_key == "easy run":
            analysis_session_type = "Easy Run"

        elif session_type_key == "long run":
            analysis_session_type = "Long Run"

        elif session_type_key in {
            "park run",
            "parkrun"
        }:
            analysis_session_type = "Park Run"

        else:
            continue

        # ------------------------------------------
        # Keep outdoor sessions only
        # ------------------------------------------

        environment = str(
            lap_record.get(
                "Environment"
            )
            or ""
        ).strip()

        if environment.lower() != "outdoor":
            continue

        # ------------------------------------------
        # Create the session grouping key
        # ------------------------------------------

        session_date = normalise_date(
            lap_record.get("Date")
        )

        session_name = str(
            lap_record.get(
                "Session Name"
            )
            or ""
        )

        session_key = (
            session_date,
            session_name,
            analysis_session_type
        )

        sessions[session_key].append(
            lap_record
        )

    drift_rows = {
        "Easy Run": [],
        "Long Run": []
    }

    work_segment_rows = []

    for (
        session_date,
        session_name,
        session_type
    ), laps in sessions.items():

        laps.sort(
            key=lambda lap: safe_float(
                lap.get("Lap Number")
            )
        )

        total_distance = sum(
            safe_float(
                lap.get(
                    "Lap Distance (km)"
                )
            )
            for lap in laps
        )

        if total_distance <= 0:
            continue

        planned_candidates = plan_by_date.get(
            session_date,
            []
        )

        plan_record = {}

        for candidate in planned_candidates:
            candidate_sport = str(
                candidate.get("Sport", "")
            ).strip().lower()

            if candidate_sport in {
                "running",
                "run"
            }:
                plan_record = candidate
                break

        drift_instruction = get_plan_value(
            plan_record,
            "Drift Segment",
            "Drift Segments"
        )

        work_instruction = get_plan_value(
            plan_record,
            "Work Segments",
            "Work Segment"
        )

        notes = get_plan_value(
            plan_record,
            "Notes"
        )

        # Park Runs are only suitable for Easy HR Drift
        # when a steady analysis segment is configured.
        if session_type == "Park Run":
            if not drift_instruction:
                continue

            output_session_type = "Easy Run"

        else:
            output_session_type = session_type

        try:
            custom_drift_segment = (
                parse_distance_segment(
                    drift_instruction
                )
            )

        except ValueError:
            custom_drift_segment = None

        if custom_drift_segment is None:
            segment_start = 0.0
            segment_end = total_distance
            analysis_source = (
                "Automatic full session"
            )

        else:
            segment_start = (
                custom_drift_segment[0]
            )

            segment_end = min(
                custom_drift_segment[1],
                total_distance
            )

            analysis_source = (
                f"Training plan: "
                f"{drift_instruction}"
            )

        segment_distance = (
            segment_end - segment_start
        )

        if segment_distance < 4:
            context_flag = (
                "Segment too short for "
                "meaningful drift"
            )

        else:
            context_flag = (
                "Review conditions, terrain "
                "and HR reliability"
            )

        drift_result = calculate_drift(
            laps,
            segment_start,
            segment_end
        )

        first_half = drift_result[
            "first_half"
        ]

        second_half = drift_result[
            "second_half"
        ]

        midpoint = drift_result[
            "midpoint"
        ]

        drift_rows[
            output_session_type
        ].append([
            session_date,
            session_name,
            (
                f"{segment_start:g}-"
                f"{segment_end:g} km"
            ),
            segment_distance,
            (
                f"{segment_start:g}-"
                f"{midpoint:g} km"
            ),
            seconds_to_excel_time(
                first_half[
                    "pace_seconds_per_km"
                ]
            ),
            first_half["average_hr"],
            first_half["ascent_m"],
            (
                f"{midpoint:g}-"
                f"{segment_end:g} km"
            ),
            seconds_to_excel_time(
                second_half[
                    "pace_seconds_per_km"
                ]
            ),
            second_half["average_hr"],
            second_half["ascent_m"],
            drift_result["hr_drift"],
            drift_result["pace_change"],
            drift_result["decoupling"],
            (
                second_half["ascent_m"]
                - first_half["ascent_m"]
            ),
            analysis_source,
            context_flag,
            notes
        ])
        
        try:
            work_segments = parse_work_segments(
                work_instruction
            )

        except ValueError as error:
            print(
                "Work segment parsing error:",
                session_date,
                session_name,
                repr(work_instruction),
                error
            )

            work_segments = []

        for work_segment in work_segments:
            work_start = work_segment[
                "start_km"
            ]

            work_end = min(
                work_segment["end_km"],
                total_distance
            )

            if work_end <= work_start:
                continue

            work_result = get_segment_from_laps(
                laps,
                work_start,
                work_end
            )

            work_segment_rows.append([
                session_date,
                session_name,
                work_segment["name"],
                work_start,
                work_end,
                work_result["distance_km"],
                seconds_to_excel_time(
                    work_result[
                        "time_seconds"
                    ]
                ),
                seconds_to_excel_time(
                    work_result[
                        "pace_seconds_per_km"
                    ]
                ),
                work_result["average_hr"],
                work_result["max_hr"],
                work_result["ascent_m"],
                work_result["efficiency"],
                notes
            ])

    # --------------------------------------------------
    # Write completed analysis to trend_summary.xlsx
    # --------------------------------------------------

    trend_workbook = load_workbook(
        TREND_FILE
    )

    for session_type, sheet_name in (
        DRIFT_SHEETS.items()
    ):
        rows = sorted(
            drift_rows[session_type],
            key=lambda row: row[0]
        )

        write_drift_sheet(
            trend_workbook,
            sheet_name,
            rows
        )

    write_work_segments_sheet(
        trend_workbook,
        sorted(
            work_segment_rows,
            key=lambda row: (
                row[0],
                row[3]
            )
        )
    )

    trend_workbook.save(
        TREND_FILE
    )

    print()
    print("=" * 50)
    print("HR DRIFT ANALYSIS COMPLETE")
    print("=" * 50)
    print(
        "Easy runs analysed:",
        len(drift_rows["Easy Run"])
    )
    print(
        "Long runs analysed:",
        len(drift_rows["Long Run"])
    )
    print(
        "Work segments analysed:",
        len(work_segment_rows)
    )
    print(
        "Updated:",
        TREND_FILE
    )


if __name__ == "__main__":
    main()