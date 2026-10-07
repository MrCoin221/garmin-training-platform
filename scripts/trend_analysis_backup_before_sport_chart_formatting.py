from openpyxl import load_workbook, Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import os

input_file = "output/session_summary.xlsx"
output_file = "output/trend_summary.xlsx"

# Open the session summary Excel file
workbook = load_workbook(input_file)
sheet = workbook["Session Summary"]

# Read header row
headers = [cell.value for cell in sheet[1]]

# Find column positions
date_col = headers.index("Date") + 1
session_name_col = headers.index("Session Name") + 1
session_type_col = headers.index("Session Type") + 1
sport_col = headers.index("Sport") + 1
environment_col = headers.index("Environment") + 1
distance_col = headers.index("Distance (km)") + 1
time_col = headers.index("Time (hh:mm:ss)") + 1
heart_rate_col = headers.index("Average Heart Rate (bpm)") + 1
ascent_col = headers.index("Total Ascent (m)") + 1
rpe_col = headers.index("RPE") + 1 if "RPE" in headers else None
calories_col = headers.index("Calories") + 1 if "Calories" in headers else None


def time_to_seconds(time_value):
    """Convert hh:mm:ss text into seconds."""
    if time_value is None:
        return 0

    time_text = str(time_value)
    parts = time_text.split(":")

    if len(parts) != 3:
        return 0

    hours = int(parts[0])
    minutes = int(parts[1])
    seconds = int(parts[2])

    return hours * 3600 + minutes * 60 + seconds


def seconds_to_hms(total_seconds):
    """Convert seconds into hh:mm:ss text."""
    total_seconds = int(total_seconds)

    hours = total_seconds // 3600
    minutes = (total_seconds % 3600) // 60
    seconds = total_seconds % 60

    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def seconds_per_km_to_pace(seconds_per_km):
    """Convert seconds per km into mm:ss pace text."""
    if seconds_per_km is None:
        return None

    minutes = int(seconds_per_km // 60)
    seconds = int(seconds_per_km % 60)

    return f"{minutes:02d}:{seconds:02d}"


def get_week_start(date_value):
    """Return Monday of the week for a given YYYY-MM-DD date."""
    if date_value is None:
        return None

    date_text = str(date_value)

    try:
        date_object = datetime.strptime(date_text, "%Y-%m-%d")
    except ValueError:
        return None

    week_start = date_object - timedelta(days=date_object.weekday())
    return week_start.strftime("%Y-%m-%d")


def create_empty_summary():
    return {
        "sessions": 0,
        "distance": 0,
        "time_seconds": 0,
        "heart_rates": [],
        "rpes": [],
        "ascent": 0
    }


def create_empty_week_summary():
    return {
        "sessions": 0,
        "distance": 0,
        "time_seconds": 0,
        "heart_rates": [],
        "rpes": [],
        "ascent": 0,
        "easy_distance": 0,
        "long_run_distance": 0,
        "quality_sessions": 0
    }


def create_empty_week_session_type_summary():
    return {
        "sessions": 0,
        "distance": 0,
        "time_seconds": 0,
        "heart_rates": [],
        "rpes": [],
        "ascent": 0
    }


def create_empty_sport_summary():
    return {
        "sessions": 0,
        "distance": 0,
        "time_seconds": 0,
        "heart_rates": [],
        "rpes": [],
        "calories": 0
    }


def calculate_efficiency(distance_km, time_seconds, heart_rates):
    """
    Calculate simple pace-HR efficiency.

    Formula:
    average speed in m/s divided by average heart rate, multiplied by 1000.

    Higher generally means more speed for the same heart rate,
    but only compare this within the same session type.
    """
    if distance_km is None or distance_km <= 0:
        return None

    if time_seconds is None or time_seconds <= 0:
        return None

    if not heart_rates:
        return None

    avg_hr = sum(heart_rates) / len(heart_rates)
    distance_m = distance_km * 1000
    avg_speed_mps = distance_m / time_seconds

    efficiency = (avg_speed_mps / avg_hr) * 1000

    return efficiency


def parse_rpe(rpe_value):
    """Convert RPE value to a number, ignoring blanks."""
    if rpe_value is None:
        return None

    try:
        return float(rpe_value)
    except ValueError:
        return None


def classify_sport_category(sport, session_name):
    sport_text = str(sport).strip().lower() if sport is not None else ""
    session_text = str(session_name).strip().lower() if session_name is not None else ""

    # Filename first
    if "cycling" in session_text or "bike" in session_text:
        return "Bike"

    if "swimming" in session_text or "swim" in session_text:
        return "Swim"

    if "run" in session_text:
        return "Run"

    # Garmin sport second
    if "cycling" in sport_text:
        return "Bike"

    if "swimming" in sport_text:
        return "Swim"

    if "running" in sport_text:
        return "Run"

    if (
        "fitness_equipment" in sport_text
        or "strength" in sport_text
        or "training" in sport_text
    ):
        return "Strength"

    return "Other"


def add_data_labels(chart):
    """Show values on chart data points."""
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True


# Running-only totals
total_sessions = 0
total_distance = 0
total_time_seconds = 0
heart_rates = []
total_ascent = 0

# All-sport context totals
all_sport_sessions = 0
all_sport_time_seconds = 0

# Running-only grouped summaries
session_type_summary = {}
environment_summary = {}
weekly_summary = {}
weekly_session_type_summary = {}

# All-sport grouped summary
weekly_sport_summary = {}

# Read data rows
for row in range(2, sheet.max_row + 1):
    date_value = sheet.cell(row=row, column=date_col).value
    session_name = sheet.cell(row=row, column=session_name_col).value
    session_type = sheet.cell(row=row, column=session_type_col).value
    sport = sheet.cell(row=row, column=sport_col).value
    sport_category = classify_sport_category(sport, session_name)
    environment = sheet.cell(row=row, column=environment_col).value
    distance = sheet.cell(row=row, column=distance_col).value
    time_value = sheet.cell(row=row, column=time_col).value
    heart_rate = sheet.cell(row=row, column=heart_rate_col).value
    ascent = sheet.cell(row=row, column=ascent_col).value

    rpe = (
        parse_rpe(sheet.cell(row=row, column=rpe_col).value)
        if rpe_col is not None
        else None
    )

    calories = (
        sheet.cell(row=row, column=calories_col).value
        if calories_col is not None
        else None
    )

    if session_type is None:
        session_type = "Unknown"

    if environment is None:
        environment = "Unknown"

    if distance is None:
        distance = 0

    if ascent is None:
        ascent = 0

    if calories is None:
        calories = 0

    time_seconds = time_to_seconds(time_value)
    week_start = get_week_start(date_value)
   

    # ---------------------------------
    # All-sport context
    # ---------------------------------
    all_sport_sessions += 1
    all_sport_time_seconds += time_seconds

    if week_start is not None:
        week_sport_key = (week_start, sport_category)

        if week_sport_key not in weekly_sport_summary:
            weekly_sport_summary[week_sport_key] = (
                create_empty_sport_summary()
            )

        sport_data = weekly_sport_summary[week_sport_key]

        sport_data["sessions"] += 1
        sport_data["distance"] += distance
        sport_data["time_seconds"] += time_seconds
        sport_data["calories"] += calories

        if heart_rate is not None:
            sport_data["heart_rates"].append(heart_rate)

        if rpe is not None:
            sport_data["rpes"].append(rpe)

    # Non-running activities stop here.
    # Bike and Swim are recorded above but cannot affect
    # running pace, HR, distance or efficiency calculations.
    if sport_category != "Run":
        continue

    # ---------------------------------
    # Running-only overall summary
    # ---------------------------------
    total_sessions += 1
    total_distance += distance
    total_time_seconds += time_seconds
    total_ascent += ascent

    if heart_rate is not None:
        heart_rates.append(heart_rate)

    # Running summary by session type
    if session_type not in session_type_summary:
        session_type_summary[session_type] = create_empty_summary()

    session_data = session_type_summary[session_type]

    session_data["sessions"] += 1
    session_data["distance"] += distance
    session_data["time_seconds"] += time_seconds
    session_data["ascent"] += ascent

    if heart_rate is not None:
        session_data["heart_rates"].append(heart_rate)

    if rpe is not None:
        session_data["rpes"].append(rpe)

    # Running summary by environment
    if environment not in environment_summary:
        environment_summary[environment] = create_empty_summary()

    environment_data = environment_summary[environment]

    environment_data["sessions"] += 1
    environment_data["distance"] += distance
    environment_data["time_seconds"] += time_seconds
    environment_data["ascent"] += ascent

    if heart_rate is not None:
        environment_data["heart_rates"].append(heart_rate)

    if rpe is not None:
        environment_data["rpes"].append(rpe)

    # Running-only weekly summary
    if week_start is not None:
        if week_start not in weekly_summary:
            weekly_summary[week_start] = create_empty_week_summary()

        week_data = weekly_summary[week_start]

        week_data["sessions"] += 1
        week_data["distance"] += distance
        week_data["time_seconds"] += time_seconds
        week_data["ascent"] += ascent

        if heart_rate is not None:
            week_data["heart_rates"].append(heart_rate)

        if rpe is not None:
            week_data["rpes"].append(rpe)

        if session_type == "Easy Run":
            week_data["easy_distance"] += distance

        if session_type == "Long Run":
            week_data["long_run_distance"] += distance

        quality_types = [
            "Threshold 1",
            "Threshold 2",
            "Tempo Run",
            "Interval Run",
            "Race Day",
            "Park Run"
        ]

        if session_type in quality_types:
            week_data["quality_sessions"] += 1

        # Running-only weekly summary by session type
        week_session_key = (week_start, session_type)

        if week_session_key not in weekly_session_type_summary:
            weekly_session_type_summary[week_session_key] = (
                create_empty_week_session_type_summary()
            )

        week_session_data = weekly_session_type_summary[week_session_key]

        week_session_data["sessions"] += 1
        week_session_data["distance"] += distance
        week_session_data["time_seconds"] += time_seconds
        week_session_data["ascent"] += ascent

        if heart_rate is not None:
            week_session_data["heart_rates"].append(heart_rate)

        if rpe is not None:
            week_session_data["rpes"].append(rpe)

# Calculate overall averages
avg_heart_rate = sum(heart_rates) / len(heart_rates) if heart_rates else None


# Print overall summary in terminal
print("\n==============================")
print("OVERALL TRAINING SUMMARY")
print("==============================")
print(f"Total sessions: {total_sessions}")
print(f"Total distance: {total_distance:.2f} km")
print(f"Total training time: {seconds_to_hms(total_time_seconds)}")
print(f"Average heart rate: {avg_heart_rate:.1f} bpm" if avg_heart_rate else "Average heart rate: None")
print(f"Total ascent: {total_ascent:.0f} m")


# Print session type summary in terminal
print("\n==============================")
print("SUMMARY BY SESSION TYPE")
print("==============================")

for session_type, data in session_type_summary.items():
    avg_hr = sum(data["heart_rates"]) / len(data["heart_rates"]) if data["heart_rates"] else None

    print(f"\n{session_type}")
    print(f"  Sessions: {data['sessions']}")
    print(f"  Distance: {data['distance']:.2f} km")
    print(f"  Time: {seconds_to_hms(data['time_seconds'])}")
    print(f"  Average HR: {avg_hr:.1f} bpm" if avg_hr else "  Average HR: None")
    print(f"  Total ascent: {data['ascent']:.0f} m")


# Print environment summary in terminal
print("\n==============================")
print("SUMMARY BY ENVIRONMENT")
print("==============================")

for environment, data in environment_summary.items():
    avg_hr = sum(data["heart_rates"]) / len(data["heart_rates"]) if data["heart_rates"] else None

    print(f"\n{environment}")
    print(f"  Sessions: {data['sessions']}")
    print(f"  Distance: {data['distance']:.2f} km")
    print(f"  Time: {seconds_to_hms(data['time_seconds'])}")
    print(f"  Average HR: {avg_hr:.1f} bpm" if avg_hr else "  Average HR: None")
    print(f"  Total ascent: {data['ascent']:.0f} m")


# Create Excel workbook for trend summary
trend_workbook = Workbook()


# -----------------------------
# Sheet 1: Overall Summary
# -----------------------------
overall_sheet = trend_workbook.active
overall_sheet.title = "Overall Summary"

overall_sheet.append(["Metric", "Value"])
overall_sheet.append(["Total Sessions", total_sessions])
overall_sheet.append(["Total Distance (km)", round(total_distance, 2)])
overall_sheet.append(["Total Training Time", seconds_to_hms(total_time_seconds)])
overall_sheet.append(["Average Heart Rate (bpm)", round(avg_heart_rate, 1) if avg_heart_rate else None])
overall_sheet.append(["Total Ascent (m)", round(total_ascent, 0)])

overall_sheet.column_dimensions["A"].width = 28
overall_sheet.column_dimensions["B"].width = 20


# -----------------------------
# Sheet 2: By Session Type
# -----------------------------
session_type_sheet = trend_workbook.create_sheet("By Session Type")

session_type_sheet.append([
    "Session Type",
    "Sessions",
    "Distance (km)",
    "Time (hh:mm:ss)",
    "Average Pace (min/km)",
    "Average Heart Rate (bpm)",
    "Total Ascent (m)",
    "Efficiency Score"
])

for session_type, data in session_type_summary.items():
    avg_hr = sum(data["heart_rates"]) / len(data["heart_rates"]) if data["heart_rates"] else None

    avg_pace = None
    if data["distance"] > 0:
        avg_pace_seconds_per_km = data["time_seconds"] / data["distance"]
        avg_pace = seconds_per_km_to_pace(avg_pace_seconds_per_km)

    efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        data["heart_rates"]
    )

    session_type_sheet.append([
        session_type,
        data["sessions"],
        round(data["distance"], 2),
        seconds_to_hms(data["time_seconds"]),
        avg_pace,
        round(avg_hr, 1) if avg_hr else None,
        round(data["ascent"], 0),
        round(efficiency, 2) if efficiency else None
    ])

session_type_widths = {
    "A": 22,
    "B": 12,
    "C": 16,
    "D": 18,
    "E": 22,
    "F": 26,
    "G": 18,
    "H": 18
}

for col, width in session_type_widths.items():
    session_type_sheet.column_dimensions[col].width = width


# -----------------------------
# Sheet 3: By Environment
# -----------------------------
environment_sheet = trend_workbook.create_sheet("By Environment")

environment_sheet.append([
    "Environment",
    "Sessions",
    "Distance (km)",
    "Time (hh:mm:ss)",
    "Average Pace (min/km)",
    "Average Heart Rate (bpm)",
    "Total Ascent (m)",
    "Efficiency Score"
])

for environment, data in environment_summary.items():
    avg_hr = sum(data["heart_rates"]) / len(data["heart_rates"]) if data["heart_rates"] else None

    avg_pace = None
    if data["distance"] > 0:
        avg_pace_seconds_per_km = data["time_seconds"] / data["distance"]
        avg_pace = seconds_per_km_to_pace(avg_pace_seconds_per_km)

    efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        data["heart_rates"]
    )

    environment_sheet.append([
        environment,
        data["sessions"],
        round(data["distance"], 2),
        seconds_to_hms(data["time_seconds"]),
        avg_pace,
        round(avg_hr, 1) if avg_hr else None,
        round(data["ascent"], 0),
        round(efficiency, 2) if efficiency else None
    ])

environment_widths = {
    "A": 18,
    "B": 12,
    "C": 16,
    "D": 18,
    "E": 22,
    "F": 26,
    "G": 18,
    "H": 18
}

for col, width in environment_widths.items():
    environment_sheet.column_dimensions[col].width = width


# -----------------------------
# Sheet 4: Weekly Summary
# -----------------------------
weekly_sheet = trend_workbook.create_sheet("Weekly Summary")

weekly_sheet.append([
    "Week Start",
    "Sessions",
    "Distance (km)",
    "Time (hh:mm:ss)",
    "Average Pace (min/km)",
    "Average Heart Rate (bpm)",
    "Average RPE",
    "RPE Entries",
    "Total Ascent (m)",
    "Easy Run Distance (km)",
    "Long Run Distance (km)",
    "Quality Sessions",
    "Efficiency Score"
])

for week_start in sorted(weekly_summary.keys()):
    data = weekly_summary[week_start]
    avg_hr = sum(data["heart_rates"]) / len(data["heart_rates"]) if data["heart_rates"] else None
    avg_rpe = sum(data["rpes"]) / len(data["rpes"]) if data["rpes"] else None

    avg_pace = None
    if data["distance"] > 0:
        avg_pace_seconds_per_km = data["time_seconds"] / data["distance"]
        avg_pace = seconds_per_km_to_pace(avg_pace_seconds_per_km)

    efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        data["heart_rates"]
    )

    weekly_sheet.append([
        week_start,
        data["sessions"],
        round(data["distance"], 2),
        seconds_to_hms(data["time_seconds"]),
        avg_pace,
        round(avg_hr, 1) if avg_hr else None,
        round(avg_rpe, 1) if avg_rpe else None,
        len(data["rpes"]),
        round(data["ascent"], 0),
        round(data["easy_distance"], 2),
        round(data["long_run_distance"], 2),
        data["quality_sessions"],
        round(efficiency, 2) if efficiency else None
    ])

weekly_widths = {
    "A": 14,
    "B": 12,
    "C": 16,
    "D": 18,
    "E": 22,
    "F": 26,
    "G": 14,
    "H": 14,
    "I": 18,
    "J": 24,
    "K": 24,
    "L": 18,
    "M": 18
}

for col, width in weekly_widths.items():
    weekly_sheet.column_dimensions[col].width = width


# -----------------------------
# Sheet 5: Weekly by Session Type
# -----------------------------
weekly_session_type_sheet = trend_workbook.create_sheet("Weekly by Session Type")

weekly_session_type_sheet.append([
    "Week Start",
    "Session Type",
    "Sessions",
    "Distance (km)",
    "Time (hh:mm:ss)",
    "Average Pace (min/km)",
    "Average Heart Rate (bpm)",
    "Average RPE",
    "RPE Entries",
    "Total Ascent (m)",
    "Efficiency Score"
])

for week_session_key in sorted(weekly_session_type_summary.keys()):
    week_start, session_type = week_session_key
    data = weekly_session_type_summary[week_session_key]

    avg_hr = sum(data["heart_rates"]) / len(data["heart_rates"]) if data["heart_rates"] else None
    avg_rpe = sum(data["rpes"]) / len(data["rpes"]) if data["rpes"] else None

    avg_pace = None
    if data["distance"] > 0:
        avg_pace_seconds_per_km = data["time_seconds"] / data["distance"]
        avg_pace = seconds_per_km_to_pace(avg_pace_seconds_per_km)

    efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        data["heart_rates"]
    )

    weekly_session_type_sheet.append([
        week_start,
        session_type,
        data["sessions"],
        round(data["distance"], 2),
        seconds_to_hms(data["time_seconds"]),
        avg_pace,
        round(avg_hr, 1) if avg_hr else None,
        round(avg_rpe, 1) if avg_rpe else None,
        len(data["rpes"]),
        round(data["ascent"], 0),
        round(efficiency, 2) if efficiency else None
    ])

weekly_session_type_widths = {
    "A": 14,
    "B": 22,
    "C": 12,
    "D": 16,
    "E": 18,
    "F": 22,
    "G": 26,
    "H": 14,
    "I": 14,
    "J": 18,
    "K": 18
}

for col, width in weekly_session_type_widths.items():
    weekly_session_type_sheet.column_dimensions[col].width = width


# -----------------------------
# Sheet 6: Easy Run Efficiency Trend
# -----------------------------
easy_efficiency_sheet = trend_workbook.create_sheet("Easy Run Efficiency Trend")

easy_efficiency_sheet.append([
    "Week Start",
    "Easy Run Sessions",
    "Distance (km)",
    "Time (hh:mm:ss)",
    "Average Pace (min/km)",
    "Average Heart Rate (bpm)",
    "Average RPE",
    "RPE Entries",
    "Total Ascent (m)",
    "Efficiency Score",
    "Change vs Previous",
    "4-Week Rolling Avg Efficiency"
])

easy_efficiency_values = []
previous_efficiency = None

for week_session_key in sorted(weekly_session_type_summary.keys()):
    week_start, session_type = week_session_key

    if session_type != "Easy Run":
        continue

    data = weekly_session_type_summary[week_session_key]

    avg_hr = sum(data["heart_rates"]) / len(data["heart_rates"]) if data["heart_rates"] else None
    avg_rpe = sum(data["rpes"]) / len(data["rpes"]) if data["rpes"] else None

    avg_pace = None
    if data["distance"] > 0:
        avg_pace_seconds_per_km = data["time_seconds"] / data["distance"]
        avg_pace = seconds_per_km_to_pace(avg_pace_seconds_per_km)

    efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        data["heart_rates"]
    )

    change_vs_previous = None
    if previous_efficiency is not None and efficiency is not None:
        change_vs_previous = efficiency - previous_efficiency

    if efficiency is not None:
        easy_efficiency_values.append(efficiency)
        previous_efficiency = efficiency

    rolling_average = None
    if easy_efficiency_values:
        last_four_values = easy_efficiency_values[-4:]
        rolling_average = sum(last_four_values) / len(last_four_values)

    easy_efficiency_sheet.append([
        week_start,
        data["sessions"],
        round(data["distance"], 2),
        seconds_to_hms(data["time_seconds"]),
        avg_pace,
        round(avg_hr, 1) if avg_hr else None,
        round(avg_rpe, 1) if avg_rpe else None,
        len(data["rpes"]),
        round(data["ascent"], 0),
        round(efficiency, 2) if efficiency else None,
        round(change_vs_previous, 2) if change_vs_previous is not None else None,
        round(rolling_average, 2) if rolling_average is not None else None
    ])

easy_efficiency_widths = {
    "A": 14,
    "B": 18,
    "C": 16,
    "D": 18,
    "E": 22,
    "F": 26,
    "G": 14,
    "H": 14,
    "I": 18,
    "J": 18,
    "K": 20,
    "L": 28
}

for col, width in easy_efficiency_widths.items():
    easy_efficiency_sheet.column_dimensions[col].width = width


 # -----------------------------
# Sheet 7: Outdoor Easy Run Efficiency Trend
# -----------------------------
outdoor_easy_sheet = trend_workbook.create_sheet("Outdoor Easy Efficiency")

outdoor_easy_sheet.append([
    "Week Start",
    "Outdoor Easy Run Sessions",
    "Distance (km)",
    "Time (hh:mm:ss)",
    "Average Pace (min/km)",
    "Average Heart Rate (bpm)",
    "Average RPE",
    "RPE Entries",
    "Total Ascent (m)",
    "Efficiency Score",
    "Change vs Previous",
    "4-Week Rolling Avg Efficiency"
])

outdoor_easy_efficiency_values = []
previous_outdoor_easy_efficiency = None

for week_session_key in sorted(weekly_session_type_summary.keys()):
    week_start, session_type = week_session_key

    if session_type != "Easy Run":
        continue

    data = weekly_session_type_summary[week_session_key]

    # We need to rebuild this from individual rows because weekly_session_type_summary
    # does not currently split Easy Runs by environment.
    outdoor_sessions = 0
    outdoor_distance = 0
    outdoor_time_seconds = 0
    outdoor_heart_rates = []
    outdoor_rpes = []
    outdoor_ascent = 0

    for row in range(2, sheet.max_row + 1):
        row_date_value = sheet.cell(row=row, column=date_col).value
        row_week_start = get_week_start(row_date_value)
        row_session_type = sheet.cell(row=row, column=session_type_col).value
        row_environment = sheet.cell(row=row, column=environment_col).value
        row_distance = sheet.cell(row=row, column=distance_col).value
        row_time_value = sheet.cell(row=row, column=time_col).value
        row_heart_rate = sheet.cell(row=row, column=heart_rate_col).value
        row_rpe = parse_rpe(sheet.cell(row=row, column=rpe_col).value) if rpe_col is not None else None
        row_ascent = sheet.cell(row=row, column=ascent_col).value

        if row_week_start != week_start:
            continue

        if row_session_type != "Easy Run":
            continue

        if row_environment != "Outdoor":
            continue

        if row_distance is None:
            row_distance = 0

        if row_ascent is None:
            row_ascent = 0

        row_time_seconds = time_to_seconds(row_time_value)

        outdoor_sessions += 1
        outdoor_distance += row_distance
        outdoor_time_seconds += row_time_seconds
        outdoor_ascent += row_ascent

        if row_heart_rate is not None:
            outdoor_heart_rates.append(row_heart_rate)

        if row_rpe is not None:
            outdoor_rpes.append(row_rpe)

    outdoor_avg_hr = (
        sum(outdoor_heart_rates) / len(outdoor_heart_rates)
        if outdoor_heart_rates
        else None
    )

    outdoor_avg_rpe = sum(outdoor_rpes) / len(outdoor_rpes) if outdoor_rpes else None

    outdoor_avg_pace = None
    if outdoor_distance > 0:
        outdoor_avg_pace_seconds_per_km = outdoor_time_seconds / outdoor_distance
        outdoor_avg_pace = seconds_per_km_to_pace(outdoor_avg_pace_seconds_per_km)

    outdoor_efficiency = calculate_efficiency(
        outdoor_distance,
        outdoor_time_seconds,
        outdoor_heart_rates
    )

    outdoor_change_vs_previous = None
    if previous_outdoor_easy_efficiency is not None and outdoor_efficiency is not None:
        outdoor_change_vs_previous = outdoor_efficiency - previous_outdoor_easy_efficiency

    if outdoor_efficiency is not None:
        outdoor_easy_efficiency_values.append(outdoor_efficiency)
        previous_outdoor_easy_efficiency = outdoor_efficiency

    outdoor_rolling_average = None
    if outdoor_easy_efficiency_values:
        last_four_values = outdoor_easy_efficiency_values[-4:]
        outdoor_rolling_average = sum(last_four_values) / len(last_four_values)

    outdoor_easy_sheet.append([
        week_start,
        outdoor_sessions,
        round(outdoor_distance, 2),
        seconds_to_hms(outdoor_time_seconds),
        outdoor_avg_pace,
        round(outdoor_avg_hr, 1) if outdoor_avg_hr else None,
        round(outdoor_avg_rpe, 1) if outdoor_avg_rpe else None,
        len(outdoor_rpes),
        round(outdoor_ascent, 0),
        round(outdoor_efficiency, 2) if outdoor_efficiency else None,
        round(outdoor_change_vs_previous, 2) if outdoor_change_vs_previous is not None else None,
        round(outdoor_rolling_average, 2) if outdoor_rolling_average is not None else None
    ])

outdoor_easy_widths = {
    "A": 14,
    "B": 24,
    "C": 16,
    "D": 18,
    "E": 22,
    "F": 26,
    "G": 14,
    "H": 14,
    "I": 18,
    "J": 18,
    "K": 20,
    "L": 28
}

for col, width in outdoor_easy_widths.items():
    outdoor_easy_sheet.column_dimensions[col].width = width


# -----------------------------
# Sheet 8: Outdoor Long Run Efficiency Trend
# -----------------------------
outdoor_long_sheet = trend_workbook.create_sheet("Outdoor Long Efficiency")

outdoor_long_sheet.append([
    "Week Start",
    "Outdoor Long Run Sessions",
    "Distance (km)",
    "Time (hh:mm:ss)",
    "Average Pace (min/km)",
    "Average Heart Rate (bpm)",
    "Average RPE",
    "RPE Entries",
    "Total Ascent (m)",
    "Efficiency Score",
    "Change vs Previous",
    "4-Week Rolling Avg Efficiency"
])

outdoor_long_efficiency_values = []
previous_outdoor_long_efficiency = None

# We build this from individual rows so we can filter by both Session Type and Environment
outdoor_long_weekly_summary = {}

for row in range(2, sheet.max_row + 1):
    row_date_value = sheet.cell(row=row, column=date_col).value
    row_week_start = get_week_start(row_date_value)
    row_session_type = sheet.cell(row=row, column=session_type_col).value
    row_environment = sheet.cell(row=row, column=environment_col).value
    row_distance = sheet.cell(row=row, column=distance_col).value
    row_time_value = sheet.cell(row=row, column=time_col).value
    row_heart_rate = sheet.cell(row=row, column=heart_rate_col).value
    row_rpe = parse_rpe(sheet.cell(row=row, column=rpe_col).value) if rpe_col is not None else None
    row_ascent = sheet.cell(row=row, column=ascent_col).value

    if row_week_start is None:
        continue

    if row_session_type != "Long Run":
        continue

    if row_environment != "Outdoor":
        continue

    if row_distance is None:
        row_distance = 0

    if row_ascent is None:
        row_ascent = 0

    row_time_seconds = time_to_seconds(row_time_value)

    if row_week_start not in outdoor_long_weekly_summary:
        outdoor_long_weekly_summary[row_week_start] = {
            "sessions": 0,
            "distance": 0,
            "time_seconds": 0,
            "heart_rates": [],
            "rpes": [],
            "ascent": 0
        }

    outdoor_long_weekly_summary[row_week_start]["sessions"] += 1
    outdoor_long_weekly_summary[row_week_start]["distance"] += row_distance
    outdoor_long_weekly_summary[row_week_start]["time_seconds"] += row_time_seconds
    outdoor_long_weekly_summary[row_week_start]["ascent"] += row_ascent

    if row_heart_rate is not None:
        outdoor_long_weekly_summary[row_week_start]["heart_rates"].append(row_heart_rate)

    if row_rpe is not None:
        outdoor_long_weekly_summary[row_week_start]["rpes"].append(row_rpe)


for week_start in sorted(outdoor_long_weekly_summary.keys()):
    data = outdoor_long_weekly_summary[week_start]

    avg_hr = (
        sum(data["heart_rates"]) / len(data["heart_rates"])
        if data["heart_rates"]
        else None
    )

    avg_rpe = sum(data["rpes"]) / len(data["rpes"]) if data["rpes"] else None

    avg_pace = None
    if data["distance"] > 0:
        avg_pace_seconds_per_km = data["time_seconds"] / data["distance"]
        avg_pace = seconds_per_km_to_pace(avg_pace_seconds_per_km)

    efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        data["heart_rates"]
    )

    change_vs_previous = None
    if previous_outdoor_long_efficiency is not None and efficiency is not None:
        change_vs_previous = efficiency - previous_outdoor_long_efficiency

    if efficiency is not None:
        outdoor_long_efficiency_values.append(efficiency)
        previous_outdoor_long_efficiency = efficiency

    rolling_average = None
    if outdoor_long_efficiency_values:
        last_four_values = outdoor_long_efficiency_values[-4:]
        rolling_average = sum(last_four_values) / len(last_four_values)

    outdoor_long_sheet.append([
        week_start,
        data["sessions"],
        round(data["distance"], 2),
        seconds_to_hms(data["time_seconds"]),
        avg_pace,
        round(avg_hr, 1) if avg_hr else None,
        round(avg_rpe, 1) if avg_rpe else None,
        len(data["rpes"]),
        round(data["ascent"], 0),
        round(efficiency, 2) if efficiency else None,
        round(change_vs_previous, 2) if change_vs_previous is not None else None,
        round(rolling_average, 2) if rolling_average is not None else None
    ])

outdoor_long_widths = {
    "A": 14,
    "B": 24,
    "C": 16,
    "D": 18,
    "E": 22,
    "F": 26,
    "G": 14,
    "H": 14,
    "I": 18,
    "J": 18,
    "K": 20,
    "L": 28
}

for col, width in outdoor_long_widths.items():
    outdoor_long_sheet.column_dimensions[col].width = width


# -----------------------------
# Helper: Outdoor Session Type Efficiency Trend
# -----------------------------
def create_outdoor_session_efficiency_sheet(sheet_name, target_session_type, sessions_header):
    efficiency_sheet = trend_workbook.create_sheet(sheet_name)

    efficiency_sheet.append([
        "Week Start",
        sessions_header,
        "Distance (km)",
        "Time (hh:mm:ss)",
        "Average Pace (min/km)",
        "Average Heart Rate (bpm)",
        "Average RPE",
        "RPE Entries",
        "Total Ascent (m)",
        "Efficiency Score",
        "Change vs Previous",
        "4-Week Rolling Avg Efficiency"
    ])

    efficiency_values = []
    previous_efficiency = None
    weekly_session_summary = {}

    for row in range(2, sheet.max_row + 1):
        row_date_value = sheet.cell(row=row, column=date_col).value
        row_week_start = get_week_start(row_date_value)
        row_session_type = sheet.cell(row=row, column=session_type_col).value
        row_environment = sheet.cell(row=row, column=environment_col).value
        row_distance = sheet.cell(row=row, column=distance_col).value
        row_time_value = sheet.cell(row=row, column=time_col).value
        row_heart_rate = sheet.cell(row=row, column=heart_rate_col).value
        row_ascent = sheet.cell(row=row, column=ascent_col).value
        row_rpe = parse_rpe(sheet.cell(row=row, column=rpe_col).value) if rpe_col is not None else None

        if row_week_start is None:
            continue

        if row_session_type != target_session_type:
            continue

        if row_environment != "Outdoor":
            continue

        if row_distance is None:
            row_distance = 0

        if row_ascent is None:
            row_ascent = 0

        row_time_seconds = time_to_seconds(row_time_value)

        if row_week_start not in weekly_session_summary:
            weekly_session_summary[row_week_start] = {
                "sessions": 0,
                "distance": 0,
                "time_seconds": 0,
                "heart_rates": [],
                "rpes": [],
                "ascent": 0
            }

        weekly_session_summary[row_week_start]["sessions"] += 1
        weekly_session_summary[row_week_start]["distance"] += row_distance
        weekly_session_summary[row_week_start]["time_seconds"] += row_time_seconds
        weekly_session_summary[row_week_start]["ascent"] += row_ascent

        if row_heart_rate is not None:
            weekly_session_summary[row_week_start]["heart_rates"].append(row_heart_rate)

        if row_rpe is not None:
            weekly_session_summary[row_week_start]["rpes"].append(row_rpe)

    for week_start in sorted(weekly_session_summary.keys()):
        data = weekly_session_summary[week_start]

        avg_hr = (
            sum(data["heart_rates"]) / len(data["heart_rates"])
            if data["heart_rates"]
            else None
        )

        avg_rpe = sum(data["rpes"]) / len(data["rpes"]) if data["rpes"] else None

        avg_pace = None
        if data["distance"] > 0:
            avg_pace_seconds_per_km = data["time_seconds"] / data["distance"]
            avg_pace = seconds_per_km_to_pace(avg_pace_seconds_per_km)

        efficiency = calculate_efficiency(
            data["distance"],
            data["time_seconds"],
            data["heart_rates"]
        )

        change_vs_previous = None
        if previous_efficiency is not None and efficiency is not None:
            change_vs_previous = efficiency - previous_efficiency

        if efficiency is not None:
            efficiency_values.append(efficiency)
            previous_efficiency = efficiency

        rolling_average = None
        if efficiency_values:
            last_four_values = efficiency_values[-4:]
            rolling_average = sum(last_four_values) / len(last_four_values)

        efficiency_sheet.append([
            week_start,
            data["sessions"],
            round(data["distance"], 2),
            seconds_to_hms(data["time_seconds"]),
            avg_pace,
            round(avg_hr, 1) if avg_hr else None,
            round(avg_rpe, 1) if avg_rpe else None,
            len(data["rpes"]),
            round(data["ascent"], 0),
            round(efficiency, 2) if efficiency else None,
            round(change_vs_previous, 2) if change_vs_previous is not None else None,
            round(rolling_average, 2) if rolling_average is not None else None
        ])


# -----------------------------
# Sheet 9: Outdoor Threshold 1 Efficiency Trend
# -----------------------------
create_outdoor_session_efficiency_sheet(
    "Outdoor Threshold 1 Efficiency",
    "Threshold 1",
    "Outdoor Threshold 1 Sessions"
)


# -----------------------------
# Sheet 10: Outdoor Tempo Efficiency Trend
# -----------------------------
create_outdoor_session_efficiency_sheet(
    "Outdoor Tempo Efficiency",
    "Tempo Run",
    "Outdoor Tempo Run Sessions"
)


# -----------------------------
# Sheet 11: Tempo Work-Block Trend
# -----------------------------
tempo_work_block_input_file = "output/tempo_work_block_summary.xlsx"

if os.path.exists(tempo_work_block_input_file):
    tempo_workbook = load_workbook(tempo_work_block_input_file)
    tempo_sheet = tempo_workbook["Tempo Work Summary"]

        # Build RPE lookup from session_summary.xlsx
    session_rpe_lookup = {}

    for row in range(2, sheet.max_row + 1):
        session_date = sheet.cell(row=row, column=date_col).value
        session_name = sheet.cell(row=row, column=session_name_col).value if "Session Name" in headers else None
        session_rpe = parse_rpe(sheet.cell(row=row, column=rpe_col).value) if rpe_col is not None else None

        if session_date is not None and session_name is not None:
            session_rpe_lookup[(session_date, session_name)] = session_rpe

    tempo_headers = [cell.value for cell in tempo_sheet[1]]

    tempo_date_col = tempo_headers.index("Date") + 1
    tempo_session_name_col = tempo_headers.index("Session Name") + 1
    tempo_work_blocks_col = tempo_headers.index("Work Blocks") + 1
    tempo_work_laps_col = tempo_headers.index("Work Laps") + 1
    tempo_work_distance_col = tempo_headers.index("Work Distance (km)") + 1
    tempo_work_time_col = tempo_headers.index("Work Time") + 1
    tempo_work_pace_col = tempo_headers.index("Work Pace") + 1
    tempo_work_avg_hr_col = tempo_headers.index("Work Avg HR") + 1
    tempo_work_max_hr_col = tempo_headers.index("Work Max HR") + 1
    tempo_work_ascent_col = tempo_headers.index("Work Ascent (m)") + 1
    tempo_work_efficiency_col = tempo_headers.index("Work Efficiency Score") + 1

    tempo_rows = []

    for row in range(2, tempo_sheet.max_row + 1):
        date = tempo_sheet.cell(row=row, column=tempo_date_col).value
        session_name = tempo_sheet.cell(row=row, column=tempo_session_name_col).value
        work_blocks = tempo_sheet.cell(row=row, column=tempo_work_blocks_col).value
        work_laps = tempo_sheet.cell(row=row, column=tempo_work_laps_col).value
        work_distance = tempo_sheet.cell(row=row, column=tempo_work_distance_col).value
        work_time = tempo_sheet.cell(row=row, column=tempo_work_time_col).value
        work_pace = tempo_sheet.cell(row=row, column=tempo_work_pace_col).value
        work_avg_hr = tempo_sheet.cell(row=row, column=tempo_work_avg_hr_col).value
        work_max_hr = tempo_sheet.cell(row=row, column=tempo_work_max_hr_col).value
        work_ascent = tempo_sheet.cell(row=row, column=tempo_work_ascent_col).value
        work_efficiency = tempo_sheet.cell(row=row, column=tempo_work_efficiency_col).value
        rpe = session_rpe_lookup.get((date, session_name))
        rpe_entries = 1 if rpe is not None else 0

        if date is None:
            continue

        tempo_rows.append({
            "date": date,
            "session_name": session_name,
            "work_blocks": work_blocks,
            "work_laps": work_laps,
            "work_distance": work_distance,
            "work_time": work_time,
            "work_pace": work_pace,
            "work_avg_hr": work_avg_hr,
            "work_max_hr": work_max_hr,
            "work_ascent": work_ascent,
            "work_efficiency": work_efficiency,
            "rpe": rpe,
            "rpe_entries": rpe_entries
        })

    tempo_rows = sorted(tempo_rows, key=lambda x: x["date"])

    previous_efficiency = None
    tempo_efficiency_values = []

    for tempo_row in tempo_rows:
        efficiency = tempo_row["work_efficiency"]

        change_vs_previous = None
        if previous_efficiency is not None and efficiency is not None:
            change_vs_previous = efficiency - previous_efficiency

        if efficiency is not None:
            tempo_efficiency_values.append(efficiency)
            previous_efficiency = efficiency

        rolling_average = None
        if tempo_efficiency_values:
            last_four_values = tempo_efficiency_values[-4:]
            rolling_average = sum(last_four_values) / len(last_four_values)

        tempo_row["change_vs_previous"] = change_vs_previous
        tempo_row["rolling_average"] = rolling_average

    tempo_trend_sheet = trend_workbook.create_sheet("Tempo Work Trend")

    tempo_trend_sheet.append([
        "Date",
        "Session Name",
        "Work Blocks",
        "Work Laps",
        "Work Distance (km)",
        "Work Time",
        "Work Pace",
        "Work Avg HR",
        "Work Max HR",
        "Average RPE",
        "RPE Entries",
        "Work Ascent (m)",
        "Work Efficiency Score",
        "Change vs Previous",
        "4-Session Rolling Avg Efficiency"
    ])

    for tempo_row in tempo_rows:
        tempo_trend_sheet.append([
            tempo_row["date"],
            tempo_row["session_name"],
            tempo_row["work_blocks"],
            tempo_row["work_laps"],
            tempo_row["work_distance"],
            tempo_row["work_time"],
            tempo_row["work_pace"],
            tempo_row["work_avg_hr"],
            tempo_row["work_max_hr"],
            tempo_row["rpe"],
            tempo_row["rpe_entries"],
            tempo_row["work_ascent"],
            tempo_row["work_efficiency"],
            round(tempo_row["change_vs_previous"], 2) if tempo_row["change_vs_previous"] is not None else None,
            round(tempo_row["rolling_average"], 2) if tempo_row["rolling_average"] is not None else None
        ])

    tempo_trend_widths = {
    "A": 14,
    "B": 18,
    "C": 14,
    "D": 12,
    "E": 18,
    "F": 14,
    "G": 14,
    "H": 14,
    "I": 14,
    "J": 14,
    "K": 14,
    "L": 16,
    "M": 20,
    "N": 20,
    "O": 30
}

    for col, width in tempo_trend_widths.items():
        tempo_trend_sheet.column_dimensions[col].width = width

else:
    print("Tempo work block summary not found. Skipping Tempo Work Trend sheet.")
    

# -----------------------------
# Sheet: Weekly by Sport
# -----------------------------
weekly_sport_sheet = trend_workbook.create_sheet("Weekly by Sport")

weekly_sport_sheet.append([
    "Week Start",
    "Sport Category",
    "Sessions",
    "Distance (km)",
    "Time (hh:mm:ss)",
    "Average Heart Rate (bpm)",
    "Average RPE",
    "RPE Entries",
    "Calories"
])

for week_sport_key in sorted(weekly_sport_summary.keys()):
    week_start, sport_category = week_sport_key
    data = weekly_sport_summary[week_sport_key]

    avg_hr = (
        sum(data["heart_rates"]) / len(data["heart_rates"])
        if data["heart_rates"]
        else None
    )

    avg_rpe = (
        sum(data["rpes"]) / len(data["rpes"])
        if data["rpes"]
        else None
    )

    weekly_sport_sheet.append([
        week_start,
        sport_category,
        data["sessions"],
        round(data["distance"], 2),
        seconds_to_hms(data["time_seconds"]),
        round(avg_hr, 1) if avg_hr is not None else None,
        round(avg_rpe, 1) if avg_rpe is not None else None,
        len(data["rpes"]),
        round(data["calories"], 0)
    ])

weekly_sport_widths = {
    "A": 14,
    "B": 18,
    "C": 12,
    "D": 16,
    "E": 18,
    "F": 26,
    "G": 14,
    "H": 14,
    "I": 14
}

for col, width in weekly_sport_widths.items():
    weekly_sport_sheet.column_dimensions[col].width = width


# -----------------------------
# Sheet 12: Dashboard V1
# -----------------------------
dashboard_sheet = trend_workbook.create_sheet("Dashboard", 0)
dashboard_sheet.sheet_view.showGridLines = False

# Basic styling
title_fill = PatternFill("solid", fgColor="1F4E78")
section_fill = PatternFill("solid", fgColor="D9EAF7")
kpi_fill = PatternFill("solid", fgColor="EAF4F8")
header_fill = PatternFill("solid", fgColor="BDD7EE")
white_font = Font(color="FFFFFF", bold=True, size=14)
title_font = Font(color="FFFFFF", bold=True, size=16)
header_font = Font(bold=True)
kpi_label_font = Font(bold=True, color="44546A")
kpi_value_font = Font(bold=True, size=14)
thin_gray = Side(style="thin", color="D9E1F2")

# Title
dashboard_sheet.merge_cells("A1:M1")
dashboard_sheet["A1"] = "Training Dashboard V1"
dashboard_sheet["A1"].fill = title_fill
dashboard_sheet["A1"].font = title_font
dashboard_sheet["A1"].alignment = Alignment(horizontal="center")

# Overall KPI section
dashboard_sheet.merge_cells("A3:M3")
dashboard_sheet["A3"] = "Executive Overview"
dashboard_sheet["A3"].fill = section_fill
dashboard_sheet["A3"].font = Font(bold=True, size=12)

overall_efficiency = calculate_efficiency(
    total_distance,
    total_time_seconds,
    heart_rates
)

latest_week = None
latest_week_data = None

if weekly_summary:
    latest_week = sorted(weekly_summary.keys())[-1]
    latest_week_data = weekly_summary[latest_week]

latest_week_distance = latest_week_data["distance"] if latest_week_data else None
latest_week_time_hours = latest_week_data["time_seconds"] / 3600 if latest_week_data else None
latest_week_avg_rpe = (
    sum(latest_week_data["rpes"]) / len(latest_week_data["rpes"])
    if latest_week_data and latest_week_data["rpes"]
    else None
)
latest_week_efficiency = (
    calculate_efficiency(
        latest_week_data["distance"],
        latest_week_data["time_seconds"],
        latest_week_data["heart_rates"]
    )
    if latest_week_data
    else None
)

kpis = [
    ("Total Sessions", total_sessions),
    ("Total Distance (km)", round(total_distance, 2)),
    ("Total Time", seconds_to_hms(total_time_seconds)),
    ("Average HR", round(avg_heart_rate, 1) if avg_heart_rate else None),
    ("Overall Efficiency", round(overall_efficiency, 2) if overall_efficiency else None),
    ("Latest Week", latest_week),
    ("Latest Week Distance", round(latest_week_distance, 2) if latest_week_distance is not None else None),
    ("Latest Week Time", seconds_to_hms(latest_week_data["time_seconds"]) if latest_week_data else None),
    ("Latest Week Avg RPE", round(latest_week_avg_rpe, 1) if latest_week_avg_rpe else None),
    ("Latest Week Efficiency", round(latest_week_efficiency, 2) if latest_week_efficiency else None),
]

kpi_start_row = 5
kpi_start_col = 1

for index, (label, value) in enumerate(kpis):
    row_offset = index // 5
    col_offset = (index % 5) * 2

    label_cell = dashboard_sheet.cell(row=kpi_start_row + row_offset * 3, column=kpi_start_col + col_offset)
    value_cell = dashboard_sheet.cell(row=kpi_start_row + row_offset * 3 + 1, column=kpi_start_col + col_offset)

    label_cell.value = label
    value_cell.value = value

    label_cell.font = kpi_label_font
    value_cell.font = kpi_value_font

    label_cell.fill = kpi_fill
    value_cell.fill = kpi_fill

    label_cell.alignment = Alignment(horizontal="center")
    value_cell.alignment = Alignment(horizontal="center")

    label_cell.border = Border(top=thin_gray, left=thin_gray, right=thin_gray)
    value_cell.border = Border(bottom=thin_gray, left=thin_gray, right=thin_gray)

    dashboard_sheet.merge_cells(
        start_row=kpi_start_row + row_offset * 3,
        start_column=kpi_start_col + col_offset,
        end_row=kpi_start_row + row_offset * 3,
        end_column=kpi_start_col + col_offset + 1
    )

    dashboard_sheet.merge_cells(
        start_row=kpi_start_row + row_offset * 3 + 1,
        start_column=kpi_start_col + col_offset,
        end_row=kpi_start_row + row_offset * 3 + 1,
        end_column=kpi_start_col + col_offset + 1
    )

# Weekly dashboard data table
weekly_data_start_row = 12

dashboard_sheet.cell(row=weekly_data_start_row, column=1).value = "Week Start"
dashboard_sheet.cell(row=weekly_data_start_row, column=2).value = "Distance (km)"
dashboard_sheet.cell(row=weekly_data_start_row, column=3).value = "Training Time (hh:mm:ss)"
dashboard_sheet.cell(row=weekly_data_start_row, column=4).value = "Sessions"
dashboard_sheet.cell(row=weekly_data_start_row, column=5).value = "Efficiency Score"
dashboard_sheet.cell(row=weekly_data_start_row, column=6).value = "Average RPE"
dashboard_sheet.cell(row=weekly_data_start_row, column=7).value = "Quality Sessions"
dashboard_sheet.cell(row=weekly_data_start_row, column=8).value = "Easy Run Distance (km)"
dashboard_sheet.cell(row=weekly_data_start_row, column=9).value = "Long Run Distance (km)"
dashboard_sheet.cell(row=weekly_data_start_row, column=10).value = "RPE Load"
dashboard_sheet.cell(row=weekly_data_start_row, column=11).value = "Training Time (hrs)"
dashboard_sheet.cell(row=weekly_data_start_row, column=12).value = "Fitness Proxy"
dashboard_sheet.cell(row=weekly_data_start_row, column=13).value = "Training Stress Proxy"

for col in range(1, 14):
    cell = dashboard_sheet.cell(row=weekly_data_start_row, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

# Build Fitness Proxy from Outdoor Easy Run Efficiency
# Fitness Proxy = 4-entry rolling average of Outdoor Easy Run Efficiency
outdoor_easy_weekly_for_dashboard = {}

for row in range(2, sheet.max_row + 1):
    row_date_value = sheet.cell(row=row, column=date_col).value
    row_week_start = get_week_start(row_date_value)
    row_session_type = sheet.cell(row=row, column=session_type_col).value
    row_environment = sheet.cell(row=row, column=environment_col).value
    row_distance = sheet.cell(row=row, column=distance_col).value
    row_time_value = sheet.cell(row=row, column=time_col).value
    row_heart_rate = sheet.cell(row=row, column=heart_rate_col).value

    if row_week_start is None:
        continue

    if row_session_type != "Easy Run":
        continue

    if row_environment != "Outdoor":
        continue

    if row_distance is None:
        row_distance = 0

    row_time_seconds = time_to_seconds(row_time_value)

    if row_week_start not in outdoor_easy_weekly_for_dashboard:
        outdoor_easy_weekly_for_dashboard[row_week_start] = {
            "distance": 0,
            "time_seconds": 0,
            "heart_rates": []
        }

    outdoor_easy_weekly_for_dashboard[row_week_start]["distance"] += row_distance
    outdoor_easy_weekly_for_dashboard[row_week_start]["time_seconds"] += row_time_seconds

    if row_heart_rate is not None:
        outdoor_easy_weekly_for_dashboard[row_week_start]["heart_rates"].append(row_heart_rate)


fitness_proxy_by_week = {}
fitness_efficiency_values = []

for week_start in sorted(weekly_summary.keys()):
    if week_start in outdoor_easy_weekly_for_dashboard:
        data = outdoor_easy_weekly_for_dashboard[week_start]

        weekly_easy_efficiency = calculate_efficiency(
            data["distance"],
            data["time_seconds"],
            data["heart_rates"]
        )

        if weekly_easy_efficiency is not None:
            fitness_efficiency_values.append(weekly_easy_efficiency)

    if fitness_efficiency_values:
        last_four_values = fitness_efficiency_values[-4:]
        fitness_proxy_by_week[week_start] = sum(last_four_values) / len(last_four_values)
    else:
        fitness_proxy_by_week[week_start] = None

current_row = weekly_data_start_row + 1

for week_start in sorted(weekly_summary.keys()):
    data = weekly_summary[week_start]

    week_avg_rpe = sum(data["rpes"]) / len(data["rpes"]) if data["rpes"] else None

    week_efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        data["heart_rates"]
    )

    week_time_hours = data["time_seconds"] / 3600
    rpe_load = week_time_hours * week_avg_rpe if week_avg_rpe is not None else None
    fitness_proxy = fitness_proxy_by_week.get(week_start)
    training_stress_proxy = rpe_load

    dashboard_sheet.cell(row=current_row, column=1).value = week_start
    dashboard_sheet.cell(row=current_row, column=2).value = round(data["distance"], 2)
    dashboard_sheet.cell(row=current_row, column=3).value = seconds_to_hms(data["time_seconds"])
    dashboard_sheet.cell(row=current_row, column=4).value = data["sessions"]
    dashboard_sheet.cell(row=current_row, column=5).value = round(week_efficiency, 2) if week_efficiency else None
    dashboard_sheet.cell(row=current_row, column=6).value = round(week_avg_rpe, 1) if week_avg_rpe else None
    dashboard_sheet.cell(row=current_row, column=7).value = data["quality_sessions"]
    dashboard_sheet.cell(row=current_row, column=8).value = round(data["easy_distance"], 2)
    dashboard_sheet.cell(row=current_row, column=9).value = round(data["long_run_distance"], 2)
    dashboard_sheet.cell(row=current_row, column=10).value = round(rpe_load, 1) if rpe_load is not None else None
    dashboard_sheet.cell(row=current_row, column=11).value = round(week_time_hours, 2)
    dashboard_sheet.cell(row=current_row, column=12).value = fitness_proxy
    dashboard_sheet.cell(row=current_row, column=13).value = training_stress_proxy

    current_row += 1

weekly_data_end_row = current_row - 1

# Chart helper references
week_categories = Reference(
    dashboard_sheet,
    min_col=1,
    min_row=weekly_data_start_row + 1,
    max_row=weekly_data_end_row
)

# Chart: Weekly Training Overview - Distance + Time
overview_chart = BarChart()
overview_chart.title = "Weekly Training Overview"
overview_chart.y_axis.title = "Distance (km)"
overview_chart.x_axis.title = "Week"

# Distance bars from column B
overview_distance_data = Reference(
    dashboard_sheet,
    min_col=2,
    min_row=weekly_data_start_row,
    max_row=weekly_data_end_row
)

overview_chart.add_data(overview_distance_data, titles_from_data=True)
overview_chart.set_categories(week_categories)

# Training time line from hidden helper column K / column 11
overview_time_chart = LineChart()

overview_time_data = Reference(
    dashboard_sheet,
    min_col=11,
    min_row=weekly_data_start_row,
    max_row=weekly_data_end_row
)

overview_time_chart.add_data(overview_time_data, titles_from_data=True)

if overview_time_chart.series:
    overview_time_chart.series[0].graphicalProperties.line.solidFill = "C00000"
    overview_time_chart.series[0].graphicalProperties.line.width = 30000

# overview_time_chart.graphical_properties.line.width = 25000

overview_time_chart.set_categories(week_categories)

# Secondary Y-axis for training time
overview_time_chart.y_axis.axId = 200
overview_time_chart.y_axis.title = "Training Time (hrs)"
overview_time_chart.y_axis.crosses = "max"
overview_time_chart.y_axis.delete = False
overview_time_chart.y_axis.tickLblPos = "nextTo"
overview_time_chart.y_axis.numFmt = "0.0"
overview_time_chart.y_axis.scaling.min = 0

# Force sensible max for weekly training hours
weekly_time_hours_values = []

for week_start in sorted(weekly_summary.keys()):
    data = weekly_summary[week_start]
    weekly_time_hours_values.append(data["time_seconds"] / 3600)

if weekly_time_hours_values:
    overview_time_chart.y_axis.scaling.max = max(weekly_time_hours_values) + 1

# Combine bar + line chart
overview_chart += overview_time_chart

overview_chart.height = 12
overview_chart.width = 24

# Primary Y-axis for distance
overview_chart.y_axis.delete = False
overview_chart.y_axis.tickLblPos = "nextTo"
overview_chart.y_axis.numFmt = "0"
overview_chart.y_axis.scaling.min = 0

overview_chart.x_axis.delete = False
overview_chart.x_axis.tickLblPos = "low"

dashboard_sheet.add_chart(overview_chart, "N12")


# Chart 2: Weekly Activity Count
sessions_chart = BarChart()
sessions_chart.title = "Weekly Activity Count"
sessions_chart.y_axis.title = "sessions"
sessions_chart.x_axis.title = "Week"
sessions_data = Reference(
    dashboard_sheet,
    min_col=4,
    min_row=weekly_data_start_row,
    max_row=weekly_data_end_row
)
sessions_chart.add_data(sessions_data, titles_from_data=True)
sessions_chart.set_categories(week_categories)
sessions_chart.height = 12
sessions_chart.width = 24
sessions_chart.y_axis.delete = False
sessions_chart.y_axis.tickLblPos = "nextTo"
sessions_chart.y_axis.numFmt = "0.0"
sessions_chart.y_axis.scaling.min = 0
sessions_chart.y_axis.majorUnit = 1
sessions_chart.x_axis.delete = False
sessions_chart.x_axis.tickLblPos = "low"
sessions_chart.legend = None
sessions_chart.legend = None
# add_data_labels(sessions_chart)
dashboard_sheet.add_chart(sessions_chart, "N36")

# Chart 3: Weekly Efficiency
efficiency_chart = LineChart()
efficiency_chart.title = "Weekly Efficiency Score"
efficiency_chart.y_axis.title = "efficiency"
efficiency_chart.x_axis.title = "Week"
efficiency_data = Reference(
    dashboard_sheet,
    min_col=5,
    min_row=weekly_data_start_row,
    max_row=weekly_data_end_row
)
efficiency_chart.add_data(efficiency_data, titles_from_data=True)
efficiency_chart.set_categories(week_categories)
efficiency_chart.height = 12
efficiency_chart.width = 24

efficiency_values = []

for week_start in sorted(weekly_summary.keys()):
    data = weekly_summary[week_start]
    week_efficiency = calculate_efficiency(
        data["distance"],
        data["time_seconds"],
        data["heart_rates"]
    )

    if week_efficiency is not None:
        efficiency_values.append(week_efficiency)

if efficiency_values:
    efficiency_chart.y_axis.scaling.min = max(0, min(efficiency_values) - 1)
    efficiency_chart.y_axis.scaling.max = max(efficiency_values) + 1

efficiency_chart.y_axis.delete = False
efficiency_chart.y_axis.tickLblPos = "nextTo"
efficiency_chart.y_axis.numFmt = "0.0"
efficiency_chart.x_axis.delete = False
efficiency_chart.x_axis.tickLblPos = "low"
efficiency_chart.legend = None

efficiency_chart.legend = None
# add_data_labels(efficiency_chart)
dashboard_sheet.add_chart(efficiency_chart, "W36")

# Chart 4: Weekly Average RPE
rpe_chart = LineChart()
rpe_chart.title = "Weekly Average RPE"
rpe_chart.y_axis.title = "RPE"
rpe_chart.x_axis.title = "Week"
rpe_data = Reference(
    dashboard_sheet,
    min_col=6,
    min_row=weekly_data_start_row,
    max_row=weekly_data_end_row
)
rpe_chart.add_data(rpe_data, titles_from_data=True)
rpe_chart.set_categories(week_categories)
rpe_chart.height = 12
rpe_chart.width = 24
rpe_chart.y_axis.delete = False
rpe_chart.y_axis.tickLblPos = "nextTo"
rpe_chart.y_axis.numFmt = "0.0"
rpe_chart.y_axis.scaling.min = 0
rpe_chart.y_axis.scaling.max = 10
rpe_chart.x_axis.delete = False
rpe_chart.x_axis.tickLblPos = "low"
rpe_chart.legend = None
rpe_chart.legend = None
# add_data_labels(rpe_chart)
dashboard_sheet.add_chart(rpe_chart, "N59")

# Session type distance pivot table for chart
session_type_pivot_start_row = weekly_data_end_row + 4
session_types = sorted(set(key[1] for key in weekly_session_type_summary.keys()))
week_starts = sorted(weekly_summary.keys())

dashboard_sheet.cell(row=session_type_pivot_start_row, column=1).value = "Week Start"

for index, session_type in enumerate(session_types):
    dashboard_sheet.cell(row=session_type_pivot_start_row, column=index + 2).value = session_type

for col in range(1, len(session_types) + 2):
    cell = dashboard_sheet.cell(row=session_type_pivot_start_row, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

pivot_row = session_type_pivot_start_row + 1

for week_start in week_starts:
    dashboard_sheet.cell(row=pivot_row, column=1).value = week_start

    for index, session_type in enumerate(session_types):
        key = (week_start, session_type)
        value = weekly_session_type_summary[key]["distance"] if key in weekly_session_type_summary else 0
        dashboard_sheet.cell(row=pivot_row, column=index + 2).value = round(value, 2)

    pivot_row += 1

pivot_end_row = pivot_row - 1
pivot_end_col = len(session_types) + 1

# ---------------------------------
# Training Time by Sport Helper Table
# ---------------------------------

sport_time_table_start_row = pivot_end_row + 4

sport_categories = ["Run", "Bike", "Swim", "Strength", "Other"]

all_weeks = sorted(
    set(
        week_start
        for week_start, sport_category in weekly_sport_summary.keys()
    )
)

dashboard_sheet.cell(
    row=sport_time_table_start_row,
    column=1
).value = "Week Start"

for index, sport_category in enumerate(sport_categories):
    dashboard_sheet.cell(
        row=sport_time_table_start_row,
        column=index + 2
    ).value = sport_category

sport_time_row = sport_time_table_start_row + 1

for week_start in all_weeks:

    dashboard_sheet.cell(
        row=sport_time_row,
        column=1
    ).value = week_start

    for index, sport_category in enumerate(sport_categories):

        key = (week_start, sport_category)

        if key in weekly_sport_summary:
            hours = (
                weekly_sport_summary[key]["time_seconds"]
                / 3600
            )
        else:
            hours = 0

        dashboard_sheet.cell(
            row=sport_time_row,
            column=index + 2
        ).value = round(hours, 2)

    sport_time_row += 1

sport_time_end_row = sport_time_row - 1

# Chart 5: Weekly RPE Load
load_chart = LineChart()
load_chart.title = "Weekly RPE Load"
load_chart.y_axis.title = "hours x RPE"
load_chart.x_axis.title = "Week"

load_data = Reference(
    dashboard_sheet,
    min_col=10,
    min_row=weekly_data_start_row,
    max_row=weekly_data_end_row
)

load_chart.add_data(load_data, titles_from_data=True)
load_chart.set_categories(week_categories)
load_chart.height = 12
load_chart.width = 24
load_chart.y_axis.delete = False
load_chart.y_axis.tickLblPos = "nextTo"
load_chart.y_axis.numFmt = "0.0"
load_chart.y_axis.scaling.min = 0
load_chart.x_axis.delete = False
load_chart.x_axis.tickLblPos = "low"
load_chart.legend = None
load_chart.legend = None
# add_data_labels(load_chart)
dashboard_sheet.add_chart(load_chart, "N82")

# Chart 6: Fitness Proxy vs Training Stress Proxy
fitness_stress_chart = BarChart()
fitness_stress_chart.title = "Fitness Proxy vs Training Stress Proxy"
fitness_stress_chart.y_axis.title = "Training Stress Proxy"
fitness_stress_chart.x_axis.title = "Week"

# Training Stress Proxy as bars from column M / column 13
stress_data = Reference(
    dashboard_sheet,
    min_col=13,
    min_row=weekly_data_start_row,
    max_row=weekly_data_end_row
)

fitness_stress_chart.add_data(stress_data, titles_from_data=True)
fitness_stress_chart.set_categories(week_categories)

# Fitness Proxy as line from column L / column 12
fitness_line_chart = LineChart()

fitness_data = Reference(
    dashboard_sheet,
    min_col=12,
    min_row=weekly_data_start_row,
    max_row=weekly_data_end_row
)

fitness_line_chart.add_data(fitness_data, titles_from_data=True)
fitness_line_chart.set_categories(week_categories)

# Secondary Y-axis for Fitness Proxy
fitness_line_chart.y_axis.axId = 300
fitness_line_chart.y_axis.title = "Fitness Proxy"
fitness_line_chart.y_axis.crosses = "max"
fitness_line_chart.y_axis.delete = False
fitness_line_chart.y_axis.tickLblPos = "nextTo"
fitness_line_chart.y_axis.numFmt = "0.0"

# Primary Y-axis for Training Stress
fitness_stress_chart.y_axis.delete = False
fitness_stress_chart.y_axis.tickLblPos = "nextTo"
fitness_stress_chart.y_axis.numFmt = "0.0"
fitness_stress_chart.y_axis.scaling.min = 0

fitness_stress_chart.x_axis.delete = False
fitness_stress_chart.x_axis.tickLblPos = "low"

fitness_stress_chart += fitness_line_chart

fitness_stress_chart.height = 12
fitness_stress_chart.width = 24

dashboard_sheet.add_chart(fitness_stress_chart, "W12")

# ---------------------------------
# Chart: Weekly Training Time by Sport
# ---------------------------------

sport_time_chart = BarChart()

sport_time_chart.type = "col"
sport_time_chart.grouping = "stacked"
sport_time_chart.overlap = 100

sport_time_chart.title = "Weekly Training Time by Sport"

sport_time_chart.y_axis.title = "Hours"
sport_time_chart.x_axis.title = "Week"

sport_time_data = Reference(
    dashboard_sheet,
    min_col=2,
    max_col=6,
    min_row=sport_time_table_start_row,
    max_row=sport_time_end_row
)

sport_time_categories = Reference(
    dashboard_sheet,
    min_col=1,
    min_row=sport_time_table_start_row + 1,
    max_row=sport_time_end_row
)

sport_time_chart.add_data(
    sport_time_data,
    titles_from_data=True
)

sport_time_chart.set_categories(
    sport_time_categories
)

sport_time_chart.height = 12
sport_time_chart.width = 28

sport_time_chart.y_axis.scaling.min = 0

dashboard_sheet.add_chart(
    sport_time_chart,
    "W82"
)

# Chart 7: Weekly Distance by Session Type
session_type_chart = BarChart()
session_type_chart.type = "col"
session_type_chart.style = 10
session_type_chart.grouping = "stacked"
session_type_chart.overlap = 100
session_type_chart.title = "Weekly Distance by Session Type"
session_type_chart.y_axis.title = "km"
session_type_chart.x_axis.title = "Week"

session_type_data = Reference(
    dashboard_sheet,
    min_col=2,
    max_col=pivot_end_col,
    min_row=session_type_pivot_start_row,
    max_row=pivot_end_row
)

session_type_categories = Reference(
    dashboard_sheet,
    min_col=1,
    min_row=session_type_pivot_start_row + 1,
    max_row=pivot_end_row
)

session_type_chart.add_data(session_type_data, titles_from_data=True)
session_type_chart.set_categories(session_type_categories)
session_type_chart.height = 14
session_type_chart.width = 28
session_type_chart.y_axis.delete = False
session_type_chart.y_axis.tickLblPos = "nextTo"
session_type_chart.y_axis.numFmt = "0.0"
session_type_chart.y_axis.scaling.min = 0
session_type_chart.x_axis.delete = False
session_type_chart.x_axis.tickLblPos = "low"
# add_data_labels(session_type_chart)
dashboard_sheet.add_chart(session_type_chart, "W59")

# Set widths
dashboard_widths = {
    "A": 14,
    "B": 16,
    "C": 22,
    "D": 12,
    "E": 18,
    "F": 14,
    "G": 16,
    "H": 22,
    "I": 22,
    "J": 14,
    "K": 14,
    "L": 18,
    "M": 22,
    "N": 16,
    "O": 16,
    "P": 16,
    "Q": 16,
    "R": 16,
    "S": 16,
    "T": 16,
    "U": 16,
    "V": 16,
    "W": 16,
    "X": 16
}

for col, width in dashboard_widths.items():
    dashboard_sheet.column_dimensions[col].width = width

dashboard_sheet.column_dimensions["K"].hidden = True

for col, width in dashboard_widths.items():
    dashboard_sheet.column_dimensions[col].width = width

dashboard_sheet.column_dimensions["K"].hidden = True


# Save trend summary workbook
trend_workbook.save(output_file)

print(f"\nTrend summary saved to: {output_file}")
